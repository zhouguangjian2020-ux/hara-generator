"""PT 场景单一数据源迁移的只读规划与校验工具。

本模块只构建内存候选结构，绝不写入 PT.json、scenarios/PT.json 或任何
审计副本。它的任务是验证：手工展开的 FUSA PT HARA 场景能在一个 PT
Domain Pack 中以 ``scenario_id`` 引用方式收敛，场景原文只保存一次。
"""

from __future__ import annotations

from copy import deepcopy
from dataclasses import dataclass
from hashlib import sha1
import re
import unicodedata
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


_REFERENCE_EVENT_ID = re.compile(r"^P_hzrd_\d{5}$")


@dataclass(frozen=True)
class PTReferenceEvent:
    """从展开版 FUSA PT HARA Sheet 读取的一条来源危害事件。"""

    reference_event_id: str
    source_row: int
    function_name: str
    reference_failure_id: str
    anomaly: str
    vehicle_hazard: str
    scenario: str
    event_description: str
    severity: int | None
    severity_reason: str | None
    exposure: int | None
    exposure_reason: str | None
    controllability: int | None
    controllability_reason: str | None
    asil: str | None
    safety_goal_id: str | None
    safety_goal: str | None
    safe_state: str | None
    ftti: str | None


def clean_text(value: Any) -> str:
    """返回适合展示的文本：保留词序和标点，仅压缩空白。"""
    if value is None:
        return ""
    text = unicodedata.normalize("NFC", str(value))
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\u00a0", " ")
    return re.sub(r"\s+", " ", text).strip()


def normalize_scenario_text(value: Any) -> str:
    """生成仅用于精确身份匹配的场景键。

    场景原文以 ``clean_text`` 的形式保存在 catalog；身份比较额外忽略所有
    空白，以避免 Excel 换行/手工排版造成同一场景被重复保存。不会删除标点、
    不改词序、不做同义词或模糊匹配，因此不同驾驶条件仍不会被合并。
    """
    return re.sub(r"\s+", "", clean_text(value))


def scenario_id_for_text(text: str) -> str:
    """根据规范化场景文本生成稳定、与导入顺序无关的 PT 场景 ID。"""
    normalized = normalize_scenario_text(text)
    if not normalized:
        raise ValueError("不能为空场景生成 scenario_id")
    digest = sha1(normalized.encode("utf-8")).hexdigest()[:12].upper()
    return f"PT_SC_{digest}"


def _as_int(value: Any) -> int | None:
    if value in (None, ""):
        return None
    try:
        return int(str(value).strip())
    except (TypeError, ValueError):
        return None


def _as_text(value: Any) -> str | None:
    text = clean_text(value)
    return text or None


def _find_hara_sheet(workbook: Any) -> Any:
    candidates = [
        sheet for sheet in workbook.worksheets
        if "HARA" in normalize_scenario_text(sheet.title).upper() and sheet.max_column >= 18
    ]
    if len(candidates) != 1:
        names = ", ".join(sheet.title for sheet in candidates) or "<none>"
        raise ValueError(f"无法唯一定位 HARA 分析 Sheet，候选: {names}")
    return candidates[0]


def extract_expanded_pt_reference_events(workbook_path: str | Path) -> list[PTReferenceEvent]:
    """读取手工展开版 PT HARA Excel 的所有 ``P_hzrd_xxxxx`` 事件。

    Excel 展示层存在分组行和合并单元格；对每个事件行，前面的功能/失效/
    异常/危害列按 Excel 的可见继承语义向下补齐。来源行号仅用于回查，不作为
    任何业务匹配主键。
    """
    path = Path(workbook_path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = _find_hara_sheet(workbook)
    inherited: list[Any] = [None] * 19
    events: list[PTReferenceEvent] = []

    for row_number, row in enumerate(
        sheet.iter_rows(min_row=3, max_row=sheet.max_row, min_col=1, max_col=19, values_only=True),
        start=3,
    ):
        values = list(row)
        for index, value in enumerate(values):
            if value not in (None, ""):
                inherited[index] = value

        event_id = normalize_scenario_text(values[0])
        if not _REFERENCE_EVENT_ID.fullmatch(event_id):
            continue

        event = PTReferenceEvent(
            reference_event_id=event_id,
            source_row=row_number,
            function_name=clean_text(inherited[1]),
            reference_failure_id=clean_text(inherited[2]),
            anomaly=clean_text(inherited[3]),
            vehicle_hazard=clean_text(inherited[4]),
            scenario=clean_text(values[5]),
            event_description=clean_text(values[6]),
            severity=_as_int(values[7]),
            severity_reason=_as_text(values[8]),
            exposure=_as_int(values[9]),
            exposure_reason=_as_text(values[10]),
            controllability=_as_int(values[11]),
            controllability_reason=_as_text(values[12]),
            asil=_as_text(values[13]),
            safety_goal_id=_as_text(values[14]),
            safety_goal=_as_text(values[15]),
            safe_state=_as_text(values[16]),
            ftti=_as_text(values[17]),
        )
        events.append(event)

    source_ids = [event.reference_event_id for event in events]
    if len(source_ids) != len(set(source_ids)):
        raise ValueError("展开版 PT HARA 中存在重复 P_hzrd 事件 ID")
    if not events:
        raise ValueError("展开版 PT HARA 中未读取到 P_hzrd 事件")
    return events


def _iter_scenario_strings(value: Any) -> Iterable[str]:
    """在旧格式的对象中收集名为 scenarios 的文本数组。"""
    if isinstance(value, dict):
        for key, child in value.items():
            if key == "scenarios" and isinstance(child, list):
                for item in child:
                    if isinstance(item, str) and normalize_scenario_text(item):
                        yield item
            else:
                yield from _iter_scenario_strings(child)
    elif isinstance(value, list):
        for child in value:
            yield from _iter_scenario_strings(child)


def _existing_scenario_texts(pack: dict[str, Any], scenario_asset: dict[str, Any]) -> dict[str, dict[str, Any]]:
    """返回 ``规范化场景键 -> 原文/来源层``，用于识别保留的 legacy 场景。"""
    result: dict[str, dict[str, Any]] = {}

    def add(text: Any, source: str) -> None:
        normalized = normalize_scenario_text(text)
        display_text = clean_text(text)
        if not normalized:
            return
        entry = result.setdefault(normalized, {"text": display_text, "sources": set()})
        entry["sources"].add(source)

    risk = pack.get("risk_catalog", {})
    catalog = risk.get("scenario_catalog", {})
    if isinstance(catalog, dict):
        for scenario_id, item in catalog.items():
            if isinstance(item, dict):
                add(item.get("text"), "domain_pack.scenario_catalog")
    for text in _iter_scenario_strings(risk.get("scenario_lookup", {})):
        add(text, "domain_pack.scenario_lookup")
    for baseline in risk.get("compatible_event_baselines", []):
        if isinstance(baseline, dict):
            add(baseline.get("scenario"), "domain_pack.compatible_event_baselines")
    for event in risk.get("event_matrix", []):
        if isinstance(event, dict):
            add(event.get("scenario"), "domain_pack.event_matrix")
    for scenario_set in risk.get("scenario_sets", {}).values():
        if not isinstance(scenario_set, dict):
            continue
        for item in scenario_set.get("scenarios", []):
            if isinstance(item, dict):
                add(item.get("text"), "domain_pack.scenario_sets")

    for text in _iter_scenario_strings(scenario_asset.get("function_types", {})):
        add(text, "scenarios/PT.json")
    return result


def _replace_lookup_scenarios_with_ids(value: Any, catalog_ids: dict[str, str]) -> Any:
    """深复制旧 lookup，并将 ``scenarios`` 文本数组替换为 ``scenario_ids``。"""
    if isinstance(value, dict):
        result: dict[str, Any] = {}
        for key, child in value.items():
            if key == "scenarios" and isinstance(child, list):
                scenario_ids: list[str] = []
                for text in child:
                    normalized = normalize_scenario_text(text)
                    if not normalized:
                        continue
                    scenario_id = catalog_ids[normalized]
                    if scenario_id not in scenario_ids:
                        scenario_ids.append(scenario_id)
                result["scenario_ids"] = scenario_ids
            else:
                result[key] = _replace_lookup_scenarios_with_ids(child, catalog_ids)
        return result
    if isinstance(value, list):
        return [_replace_lookup_scenarios_with_ids(child, catalog_ids) for child in value]
    return deepcopy(value)


def _merge_lookup_with_legacy_asset(pack_lookup: dict[str, Any], scenario_asset: dict[str, Any]) -> tuple[dict[str, Any], list[str]]:
    """按路径合并当前 PT Pack lookup 与 scenarios/PT.json，记录非场景字段冲突。

    只合并关键词、异常关键词、来源异常和场景列表；其他同名非空字段发生不同时
    不擅自选边，返回冲突路径供实际迁入前人工确认。
    """
    merged = deepcopy(pack_lookup)
    conflicts: list[str] = []

    for function_type, asset_type in scenario_asset.get("function_types", {}).items():
        if not isinstance(asset_type, dict):
            continue
        target_type = merged.setdefault(function_type, {})
        if not isinstance(target_type, dict):
            conflicts.append(f"scenario_lookup.{function_type}: Pack 类型不是对象")
            continue
        for key in ("keywords",):
            values = asset_type.get(key, [])
            if isinstance(values, list):
                target = target_type.setdefault(key, [])
                if not isinstance(target, list):
                    conflicts.append(f"scenario_lookup.{function_type}.{key}: Pack 类型不是数组")
                else:
                    for item in values:
                        if item not in target:
                            target.append(item)

        asset_modes = asset_type.get("failure_modes", {})
        target_modes = target_type.setdefault("failure_modes", {})
        if not isinstance(asset_modes, dict) or not isinstance(target_modes, dict):
            conflicts.append(f"scenario_lookup.{function_type}.failure_modes: 类型不兼容")
            continue
        for failure_mode, asset_mode in asset_modes.items():
            if not isinstance(asset_mode, dict):
                continue
            target_mode = target_modes.setdefault(failure_mode, {})
            if not isinstance(target_mode, dict):
                conflicts.append(f"scenario_lookup.{function_type}.failure_modes.{failure_mode}: Pack 类型不是对象")
                continue
            for key in ("anomaly_keywords", "source_anomalies", "scenarios"):
                values = asset_mode.get(key, [])
                if not isinstance(values, list):
                    continue
                target = target_mode.setdefault(key, [])
                if not isinstance(target, list):
                    conflicts.append(f"scenario_lookup.{function_type}.{failure_mode}.{key}: Pack 类型不是数组")
                    continue
                for item in values:
                    if item not in target:
                        target.append(item)
            for key in ("safety_goal", "safe_state", "ftti", "guidance"):
                source_value = asset_mode.get(key)
                target_value = target_mode.get(key)
                if source_value in (None, ""):
                    continue
                if target_value in (None, ""):
                    target_mode[key] = source_value
                elif normalize_scenario_text(source_value) != normalize_scenario_text(target_value):
                    conflicts.append(
                        f"scenario_lookup.{function_type}.{failure_mode}.{key}: PT Pack 与 scenarios/PT.json 不一致"
                    )
    return merged, sorted(set(conflicts))


def _reference_event_record(event: PTReferenceEvent, scenario_id: str) -> dict[str, Any]:
    """构造候选 Pack 内的来源事件记录；场景文本不重复保存。"""
    return {
        "reference_source": "FUSA_PT_EXPANDED_2026_08_25",
        "reference_event_id": event.reference_event_id,
        "source_row": event.source_row,
        "function_name": event.function_name,
        "reference_failure_id": event.reference_failure_id,
        "anomaly": event.anomaly,
        "vehicle_hazard": event.vehicle_hazard,
        "scenario_id": scenario_id,
        "description": event.event_description or None,
        "severity": event.severity,
        "severity_reason": event.severity_reason,
        "exposure": event.exposure,
        "exposure_reason": event.exposure_reason,
        "controllability": event.controllability,
        "controllability_reason": event.controllability_reason,
        "asil": event.asil,
        "safety_goal_id": event.safety_goal_id,
        "safety_goal": event.safety_goal,
        "safe_state": event.safe_state,
        "ftti": event.ftti,
    }


def build_pt_single_source_candidate(
    pack: dict[str, Any],
    scenario_asset: dict[str, Any],
    source_events: list[PTReferenceEvent],
) -> tuple[dict[str, Any], dict[str, Any]]:
    """构建 PT 单源候选对象与统计摘要，不修改调用方对象、不写文件。

    该候选刻意不把任何 FUSA 单项目结论提升为 exact/compatible；它只验证场景
    文本去重、引用完整性和来源事件收敛。事件级证据状态将在正式迁入评审中按
    两项目交集/项目变体/needs_review 决定。
    """
    if not isinstance(pack, dict):
        raise ValueError("PT Pack 必须为对象")
    if not source_events:
        raise ValueError("source_events 不能为空")

    risk = pack.get("risk_catalog", {})
    if not isinstance(risk, dict):
        raise ValueError("PT Pack 缺少 risk_catalog")

    existing = _existing_scenario_texts(pack, scenario_asset)
    source_by_key: dict[str, list[PTReferenceEvent]] = {}
    source_display_text: dict[str, str] = {}
    for event in source_events:
        scenario_key = normalize_scenario_text(event.scenario)
        if not scenario_key:
            raise ValueError(f"{event.reference_event_id} 缺少展开场景")
        source_by_key.setdefault(scenario_key, []).append(event)
        source_display_text.setdefault(scenario_key, event.scenario)

    # 来源展开场景优先作为展示原文；仅历史存在的两类场景保留为 legacy 待确认。
    all_keys = set(source_by_key) | set(existing)
    catalog_ids = {key: scenario_id_for_text(key) for key in all_keys}
    reverse_ids: dict[str, str] = {}
    for key, scenario_id in catalog_ids.items():
        old_key = reverse_ids.setdefault(scenario_id, key)
        if old_key != key:
            raise ValueError(f"scenario_id 哈希碰撞: {scenario_id}")

    catalog: dict[str, dict[str, Any]] = {}
    for key in sorted(all_keys, key=lambda item: catalog_ids[item]):
        scenario_id = catalog_ids[key]
        catalog[scenario_id] = {
            "text": source_display_text[key] if key in source_display_text else existing[key]["text"],
            "lifecycle_status": "active_reference" if key in source_by_key else "legacy_pending_confirmation",
            "source_kinds": sorted(
                ({"FUSA_PT_EXPANDED_2026_08_25"} if key in source_by_key else set())
                | existing.get(key, {}).get("sources", set())
            ),
        }

    merged_lookup, lookup_conflicts = _merge_lookup_with_legacy_asset(
        risk.get("scenario_lookup", {}), scenario_asset
    )
    candidate = deepcopy(pack)
    candidate_risk = candidate.setdefault("risk_catalog", {})
    candidate_risk["scenario_catalog"] = catalog
    candidate_risk["scenario_lookup"] = _replace_lookup_scenarios_with_ids(merged_lookup, catalog_ids)

    baselines: list[dict[str, Any]] = []
    for baseline in candidate_risk.get("compatible_event_baselines", []):
        if not isinstance(baseline, dict):
            raise ValueError("compatible_event_baselines 包含非对象记录")
        rewritten = deepcopy(baseline)
        scenario = normalize_scenario_text(rewritten.pop("scenario", ""))
        existing_scenario_id = rewritten.get("scenario_id")
        if scenario:
            rewritten["scenario_id"] = catalog_ids[scenario]
        elif isinstance(existing_scenario_id, str) and existing_scenario_id in catalog:
            # 已发布单源 Pack 的再次 dry-run：保留稳定 ID，不退回文本字段。
            rewritten["scenario_id"] = existing_scenario_id
        else:
            raise ValueError("compatible_event_baselines 存在空场景或无效 scenario_id")
        baselines.append(rewritten)
    candidate_risk["compatible_event_baselines"] = baselines

    # reference_event_matrix 是来源事件真值的候选承载位置。它不存 scenario 文本，
    # 因此与 lookup/baseline 共用唯一 scenario_catalog，避免再次复制场景列。
    candidate_risk["reference_event_matrix"] = [
        _reference_event_record(event, catalog_ids[normalize_scenario_text(event.scenario)])
        for event in source_events
    ]
    candidate_risk["scenario_single_source_contract"] = {
        "catalog_key": "risk_catalog.scenario_catalog",
        "reference_source": "FUSA_PT_EXPANDED_2026_08_25",
        "migration_state": "dry_run_candidate_not_published",
        "external_pt_scenario_file": "references/scenarios/PT.json",
        "external_pt_scenario_file_action": "remove_after_runtime_switch_and_full_regression",
    }

    source_scene_ids = {catalog_ids[key] for key in source_by_key}
    existing_scene_ids = {catalog_ids[key] for key in existing}
    summary = {
        "source_event_count": len(source_events),
        "source_unique_scenarios": len(source_by_key),
        "existing_pack_or_asset_unique_scenarios": len(existing),
        "existing_scene_texts_reused": len(set(source_by_key) & set(existing)),
        "source_scene_texts_not_in_existing_assets": len(set(source_by_key) - set(existing)),
        "legacy_scene_texts_not_in_expanded_source": len(set(existing) - set(source_by_key)),
        "candidate_catalog_count": len(catalog),
        "candidate_reference_event_count": len(candidate_risk["reference_event_matrix"]),
        "candidate_compatible_baseline_count": len(baselines),
        "source_scene_ids": source_scene_ids,
        "existing_scene_ids": existing_scene_ids,
        "lookup_conflicts": lookup_conflicts,
    }
    return candidate, summary


def validate_pt_single_source_candidate(candidate: dict[str, Any]) -> list[str]:
    """验证内存候选：场景原文只有 catalog 一份，所有引用均有效。"""
    errors: list[str] = []
    risk = candidate.get("risk_catalog", {})
    if not isinstance(risk, dict):
        return ["候选缺少 risk_catalog"]
    catalog = risk.get("scenario_catalog")
    if not isinstance(catalog, dict) or not catalog:
        return ["候选缺少非空 scenario_catalog"]

    normalized_texts: dict[str, str] = {}
    for scenario_id, item in catalog.items():
        if not isinstance(scenario_id, str) or not scenario_id.startswith("PT_SC_"):
            errors.append(f"scenario_catalog 含无效 ID: {scenario_id!r}")
            continue
        if not isinstance(item, dict):
            errors.append(f"scenario_catalog.{scenario_id} 必须是对象")
            continue
        text = normalize_scenario_text(item.get("text"))
        if not text:
            errors.append(f"scenario_catalog.{scenario_id} 缺少 text")
            continue
        if scenario_id_for_text(text) != scenario_id:
            errors.append(f"scenario_catalog.{scenario_id} 与 text 的稳定 ID 不一致")
        prior = normalized_texts.setdefault(text, scenario_id)
        if prior != scenario_id:
            errors.append(f"场景文本重复存储: {prior} / {scenario_id}")

    def assert_ids(ids: Any, path: str) -> None:
        if not isinstance(ids, list):
            errors.append(f"{path} 必须是 scenario_ids 数组")
            return
        if len(ids) != len(set(ids)):
            errors.append(f"{path} 含重复 scenario_id")
        for scenario_id in ids:
            if scenario_id not in catalog:
                errors.append(f"{path} 引用未知 scenario_id: {scenario_id}")

    def walk_lookup(value: Any, path: str = "risk_catalog.scenario_lookup") -> None:
        if isinstance(value, dict):
            if "scenarios" in value:
                errors.append(f"{path} 仍保存 scenarios 文本数组")
            if "scenario_ids" in value:
                assert_ids(value["scenario_ids"], f"{path}.scenario_ids")
            for key, child in value.items():
                walk_lookup(child, f"{path}.{key}")
        elif isinstance(value, list):
            for index, child in enumerate(value):
                walk_lookup(child, f"{path}[{index}]")

    walk_lookup(risk.get("scenario_lookup", {}))

    for index, baseline in enumerate(risk.get("compatible_event_baselines", [])):
        path = f"compatible_event_baselines[{index}]"
        if not isinstance(baseline, dict):
            errors.append(f"{path} 必须是对象")
            continue
        if "scenario" in baseline:
            errors.append(f"{path} 仍保存重复 scenario 文本")
        scenario_id = baseline.get("scenario_id")
        if scenario_id not in catalog:
            errors.append(f"{path} 引用未知 scenario_id")

    source_event_ids: set[str] = set()
    for index, event in enumerate(risk.get("reference_event_matrix", [])):
        path = f"reference_event_matrix[{index}]"
        if not isinstance(event, dict):
            errors.append(f"{path} 必须是对象")
            continue
        if "scenario" in event:
            errors.append(f"{path} 不得重复保存 scenario 文本")
        event_id = event.get("reference_event_id")
        if not isinstance(event_id, str) or not _REFERENCE_EVENT_ID.fullmatch(event_id):
            errors.append(f"{path} reference_event_id 无效")
        elif event_id in source_event_ids:
            errors.append(f"{path} reference_event_id 重复: {event_id}")
        else:
            source_event_ids.add(event_id)
        scenario_id = event.get("scenario_id")
        if scenario_id not in catalog:
            errors.append(f"{path} 引用未知 scenario_id")
        for field in ("function_name", "reference_failure_id", "anomaly", "vehicle_hazard"):
            if not normalize_scenario_text(event.get(field)):
                errors.append(f"{path} 缺少 {field}")
    return errors
