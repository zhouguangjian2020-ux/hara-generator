"""PT 扁平案例表导入与规范化工具。

该模块只服务 PT 离线参考资产导入，不参与其他域运行时逻辑。来源 Excel 的
项目 ID 仅作为 trace 保存，绝不作为跨项目语义主键。
"""
from __future__ import annotations

import hashlib
import json
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any

# 允许直接从 scripts/tools 调用时导入现有 PT ASIL 计算实现。
_SCRIPT_DIR = Path(__file__).resolve().parent.parent
if str(_SCRIPT_DIR) not in sys.path:
    sys.path.insert(0, str(_SCRIPT_DIR))

from utils.hara_rules import compute_asil  # noqa: E402
from utils.function_matcher import resolve_pt_function_semantics  # noqa: E402


CANONICAL_HEADERS = {
    "功能id": "function_id",
    "功能名称": "function_name",
    "失效id": "failure_id",
    "失效模式": "failure_mode",
    "危害事件id": "hazard_event_id",
    "hara功能名称": "hara_function_name",
    "功能异常表现": "anomaly",
    "整车危害": "vehicle_hazard",
    "运行场景": "scenario",
    "危害事件描述": "event_description",
    "s": "severity",
    "s理由": "severity_reason",
    "e": "exposure",
    "e理由": "exposure_reason",
    "c": "controllability",
    "c理由": "controllability_reason",
    "asil": "source_asil",
    "安全目标id": "safety_goal_id",
    "整车安全目标id": "vehicle_safety_goal_id",
    "安全目标": "safety_goal",
    "安全状态": "safe_state",
    "ftti": "ftti",
    "注释": "note",
    "数据来源": "source_project",
    "填写人": "filled_by",
    "填写时间": "filled_at",
}

FAILURE_MODE_ALIASES = {
    "过大": "过多",
    "过小": "过少",
}

NO_HAZARD_TOKENS = ("不涉及", "无危害", "无整车危害", "无整车层面危害")

# 这是离线导入的 PT 专属最小语义映射。后续 M3 会将它迁移为独立 alias 资产，
# 这里保留显式规则，避免按来源 func_id 猜测功能。
FUNCTION_RULES: list[dict[str, Any]] = [
    {
        "family": "pt_gear_state_control",
        "names": ("档位控制及显示功能", "档位控制及显示", "档位控制"),
        "roles": {"d档控制": "gear_shift_control", "n档控制": "gear_shift_control", "r档控制": "gear_shift_control", "p档控制": "gear_shift_control", "档位显示": "gear_information_display"},
        "disposition": "analyze",
    },
    {
        "family": "pt_hv_power_state_management",
        "names": ("整车高压上下电管理",),
        "roles": {},
        "disposition": "analyze",
    },
    {
        "family": "pt_thermal_management",
        "names": ("车辆热管理功能", "整车热管理功能"),
        "roles": {},
        "disposition": "analyze",
    },
    {
        "family": "pt_traction_torque_control",
        "names": ("车辆驱动控制", "车辆扭矩控制", "能量回收功能", "制动能量回收"),
        "roles": {"制动能量回收": "braking_regeneration", "滑行能量回收": "coasting_regeneration", "提供要求的驱动扭矩": "traction_torque_calculation"},
        "disposition": "analyze",
    },
    {
        "family": "pt_charge_discharge",
        "names": ("充放电功能", "充放电控制"),
        "roles": {},
        "disposition": "analyze",
    },
    {
        "family": "pt_hv_safety",
        "names": ("高压安全",),
        "roles": {},
        "disposition": "analyze",
    },
    {
        "family": "pt_energy_range_display",
        "names": ("车辆能耗及续航显示功能", "续驶里程和电耗显示"),
        "roles": {"default": "display_only"},
        "disposition": "exclude",
    },
]


def normalize_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    text = unicodedata.normalize("NFC", str(value)).replace("\u3000", " ").replace("\xa0", " ")
    text = re.sub(r"[\r\n\t]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def header_key(value: Any) -> str:
    return re.sub(r"[\s_]+", "", normalize_text(value)).lower()


def parse_optional_int(value: Any, *, minimum: int, maximum: int) -> int | None:
    text = normalize_text(value)
    if not text or text in {"-", "—", "null", "none", "不涉及"}:
        return None
    try:
        number = int(float(text))
    except (TypeError, ValueError):
        return None
    return number if minimum <= number <= maximum else number


def normalize_failure_mode(value: Any) -> tuple[str, str]:
    raw = normalize_text(value)
    return raw, FAILURE_MODE_ALIASES.get(raw, raw)


def classify_function(function_name: str, hara_function_name: str, pt_pack: dict[str, Any] | None = None) -> dict[str, Any]:
    name = normalize_text(function_name)
    hara_name = normalize_text(hara_function_name)
    if pt_pack is not None:
        resolved = resolve_pt_function_semantics(name, pt_pack, [hara_name])
        if resolved.get("status") in {"resolved", "unknown", "ambiguous", "needs_review"}:
            return {
                "canonical_function_family": resolved.get("canonical_function_family"),
                "semantic_role": resolved.get("semantic_role"),
                "disposition": resolved.get("disposition", "needs_review"),
                "function_match_status": resolved.get("status"),
            }
    for rule in FUNCTION_RULES:
        if name in rule["names"]:
            role = ""
            for marker, candidate in rule["roles"].items():
                if marker != "default" and marker in hara_name:
                    role = candidate
                    break
            if not role:
                role = rule["roles"].get("default", "")
            return {
                "canonical_function_family": rule["family"],
                "semantic_role": role or None,
                "disposition": rule["disposition"],
                "function_match_status": "exact_alias",
            }
    if name == "动力防盗功能":
        return {
            "canonical_function_family": None,
            "semantic_role": None,
            "disposition": "needs_review",
            "function_match_status": "unknown",
        }
    return {
        "canonical_function_family": None,
        "semantic_role": None,
        "disposition": "needs_review",
        "function_match_status": "unknown",
    }


def scenario_profile(function_name: str, hara_function_name: str, scenario: str) -> dict[str, Any]:
    text = normalize_text(scenario)
    combined = " ".join(x for x in (function_name, hara_function_name, text) if x)
    if "充电" in combined:
        maneuver = "充电"
    elif "放电" in combined or "V2L" in combined:
        maneuver = "放电"
    elif "制动" in combined or "刹车" in combined:
        maneuver = "制动"
    elif "起步" in combined or "启动" in combined:
        maneuver = "起步/启动"
    elif "停车" in combined or "泊车" in combined:
        maneuver = "停车/泊车"
    elif "换挡" in combined or "档" in combined or "挡" in combined:
        maneuver = "档位操作"
    else:
        maneuver = "未识别"

    speed_band = "未识别"
    for pattern, label in ((r"0\s*[<＜].*?10\s*km/h", "0-10km/h"), (r"10\s*[<＜].*?30\s*km/h", "10-30km/h"), (r"30\s*[<＜].*?60\s*km/h", "30-60km/h"), (r"60\s*[<＜].*?120\s*km/h", "60-120km/h")):
        if re.search(pattern, combined, re.I):
            speed_band = label
            break
    road_type = "高速公路" if "高速公路" in combined else "城市道路" if "城市" in combined else "乡村道路" if "乡村" in combined else "停车场" if "停车场" in combined else "普通道路" if "道路" in combined else "未识别"
    if "行人" in combined or "非机动车" in combined:
        traffic_condition = "含弱势道路使用者"
    elif "前方车辆" in combined or "后方车辆" in combined or "后车" in combined:
        traffic_condition = "含其他车辆"
    else:
        traffic_condition = "未识别"
    profile = {
        "maneuver": maneuver,
        "road_type": road_type,
        "speed_band": speed_band,
        "traffic_condition": traffic_condition,
        "following_distance": "较近" if "较近" in combined or "距离近" in combined else None,
        "slope": "坡道" if "坡" in combined else None,
        "environment": "低能见度/恶劣天气" if any(x in combined for x in ("雨", "雪", "雾", "冰")) else None,
        "vulnerable_road_user": True if "行人" in combined or "非机动车" in combined else False,
    }
    return {"dimensions": profile, "scenario_text": text}


def scenario_fingerprint(profile: dict[str, Any]) -> str:
    payload = json.dumps(profile, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return "PT_SC_" + hashlib.sha256(payload.encode("utf-8")).hexdigest()[:12].upper()


def _hazard_lookup(hazard_catalog: dict[str, Any]) -> dict[str, tuple[str, str]]:
    result: dict[str, tuple[str, str]] = {}
    for hazard_id, item in hazard_catalog.items():
        if not isinstance(item, dict):
            continue
        text = normalize_text(item.get("vehicle_hazard"))
        if text:
            result[text] = (hazard_id, normalize_text(item.get("hazard_family")))
    return result


def normalize_row(
    raw: dict[str, Any],
    *,
    source_file: str,
    source_sheet: str,
    source_row: int,
    hazard_catalog: dict[str, Any],
    pt_pack: dict[str, Any] | None = None,
) -> dict[str, Any]:
    function_name = normalize_text(raw.get("function_name"))
    hara_function_name = normalize_text(raw.get("hara_function_name"))
    vehicle_hazard = normalize_text(raw.get("vehicle_hazard"))
    hazard_event_id = normalize_text(raw.get("hazard_event_id"))
    vehicle_safety_goal_id = normalize_text(raw.get("vehicle_safety_goal_id")) or None
    scenario = normalize_text(raw.get("scenario"))
    failure_mode_raw, failure_mode = normalize_failure_mode(raw.get("failure_mode"))
    function = classify_function(function_name, hara_function_name, pt_pack)
    no_hazard = hazard_event_id in NO_HAZARD_TOKENS or any(token in vehicle_hazard for token in NO_HAZARD_TOKENS) or "无危害" in normalize_text(raw.get("event_description"))
    hazards = _hazard_lookup(hazard_catalog)
    hazard_id, hazard_family = hazards.get(vehicle_hazard, (None, None))
    if no_hazard:
        hazard_mapping_status = "not_applicable"
        hazard_id = None
        hazard_family = None
    elif hazard_id:
        hazard_mapping_status = "exact_catalog_match"
    else:
        hazard_mapping_status = "unresolved"

    severity = parse_optional_int(raw.get("severity"), minimum=0, maximum=3)
    exposure = parse_optional_int(raw.get("exposure"), minimum=1, maximum=4)
    controllability = parse_optional_int(raw.get("controllability"), minimum=0, maximum=3)
    computed_asil = None
    asil_status = "not_computed"
    if severity == 0:
        computed_asil = None
        asil_status = "s0_no_asil"
    elif severity is not None and exposure is not None and controllability is not None:
        try:
            computed_asil, _ = compute_asil(severity, exposure, controllability)
            asil_status = "computed"
        except ValueError:
            asil_status = "invalid_sec_combination"
    elif not no_hazard:
        asil_status = "missing_sec"

    profile = scenario_profile(function_name, hara_function_name, scenario)
    semantic_payload = {
        "canonical_function_family": function["canonical_function_family"],
        "semantic_role": function["semantic_role"],
        "failure_mode": failure_mode,
        "anomaly": normalize_text(raw.get("anomaly")),
        "hazard_family": hazard_family,
        "vehicle_hazard": vehicle_hazard if not no_hazard else None,
        "scenario_fingerprint": scenario_fingerprint(profile) if scenario else None,
    }
    semantic_key = hashlib.sha256(json.dumps(semantic_payload, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")).hexdigest()[:24]
    source_asil = normalize_text(raw.get("source_asil")) or None
    record_status = "valid_no_hazard" if no_hazard else "needs_review"
    if function["function_match_status"] == "unknown" or hazard_mapping_status == "unresolved":
        record_status = "needs_review"
    if asil_status == "invalid_sec_combination":
        record_status = "conflict"

    return {
        "record_kind": "pt_reference_case",
        "semantic_key": semantic_key,
        "canonical_function_family": function["canonical_function_family"],
        "semantic_role": function["semantic_role"],
        "disposition": function["disposition"],
        "function_match_status": function["function_match_status"],
        "function_name": function_name or None,
        "hara_function_name": hara_function_name or None,
        "normalized_failure_mode": failure_mode,
        "source_failure_mode": failure_mode_raw,
        "anomaly": normalize_text(raw.get("anomaly")),
        "vehicle_hazard": vehicle_hazard or None,
        "vehicle_hazard_id": hazard_id,
        "hazard_family": hazard_family,
        # This is an explicit reference mapping from the reviewed Excel source,
        # not a semantic matching key. It is used only to prefill S5 vehicle SGs.
        "vehicle_safety_goal_id": vehicle_safety_goal_id,
        "hazard_mapping_status": hazard_mapping_status,
        "scenario_id": profile["scenario_text"] and scenario_fingerprint(profile),
        "scenario_profile": profile,
        "description": normalize_text(raw.get("event_description")),
        "severity": severity,
        "severity_reason": normalize_text(raw.get("severity_reason")) or None,
        "exposure": None if severity == 0 else exposure,
        "exposure_reason": None if severity == 0 else (normalize_text(raw.get("exposure_reason")) or None),
        "controllability": None if severity == 0 else controllability,
        "controllability_reason": None if severity == 0 else (normalize_text(raw.get("controllability_reason")) or None),
        "source_expected_asil": source_asil,
        "computed_asil": computed_asil,
        "asil_status": asil_status,
        "safety_goal": None if severity == 0 or no_hazard else (normalize_text(raw.get("safety_goal")) or None),
        "safe_state": None if severity == 0 or no_hazard else (normalize_text(raw.get("safe_state")) or None),
        "ftti": None if severity == 0 or no_hazard else (normalize_text(raw.get("ftti")) or None),
        "evidence_status": "reference_candidate" if not no_hazard else "valid_no_hazard",
        "record_status": record_status,
        "trace": {
            "source_file": source_file,
            "source_sheet": source_sheet,
            "source_row": source_row,
            "source_project": normalize_text(raw.get("source_project")) or None,
            "raw_function_id": normalize_text(raw.get("function_id")) or None,
            "raw_failure_id": normalize_text(raw.get("failure_id")) or None,
            "raw_hazard_event_id": hazard_event_id or None,
            "raw_safety_goal_id": normalize_text(raw.get("safety_goal_id")) or None,
            "raw_vehicle_safety_goal_id": vehicle_safety_goal_id,
            "filled_by": normalize_text(raw.get("filled_by")) or None,
            "filled_at": normalize_text(raw.get("filled_at")) or None,
        },
    }


def _make_headers(ws) -> dict[str, int]:
    headers: dict[str, int] = {}
    for col in range(1, ws.max_column + 1):
        raw = ws.cell(1, col).value
        key = header_key(raw)
        if key in CANONICAL_HEADERS:
            headers[CANONICAL_HEADERS[key]] = col
    return headers


def _raw_row(ws, headers: dict[str, int], row: int) -> dict[str, Any]:
    return {name: ws.cell(row, col).value for name, col in headers.items()}


def audit_records(records: list[dict[str, Any]]) -> dict[str, Any]:
    def count_status(field: str) -> dict[str, int]:
        return dict(Counter(str(record.get(field)) for record in records))

    raw_failure: dict[str, set[str]] = defaultdict(set)
    raw_failure_rows: dict[str, list[int]] = defaultdict(list)
    raw_hazard: dict[str, set[str]] = defaultdict(set)
    raw_sg: dict[str, set[str]] = defaultdict(set)
    for record in records:
        trace = record["trace"]
        failure_id = trace.get("raw_failure_id")
        if failure_id:
            raw_failure[failure_id].add(record.get("canonical_function_family") or "unknown")
            raw_failure_rows[failure_id].append(trace["source_row"])
        hazard_id = trace.get("raw_hazard_event_id")
        if hazard_id:
            raw_hazard[hazard_id].add(record.get("semantic_key"))
        sg_id = trace.get("raw_safety_goal_id")
        if sg_id:
            raw_sg[sg_id].add(record.get("semantic_key"))

    function_counts = Counter(record["trace"].get("raw_function_id") or "<missing>" for record in records)
    source_counts = Counter(record["trace"].get("source_project") or "<missing>" for record in records)
    return {
        "record_count": len(records),
        "source_project_counts": dict(source_counts),
        "source_function_id_counts": dict(function_counts),
        "function_match_status": count_status("function_match_status"),
        "record_status": count_status("record_status"),
        "hazard_mapping_status": count_status("hazard_mapping_status"),
        "asil_status": count_status("asil_status"),
        "evidence_status": count_status("evidence_status"),
        "raw_failure_id_duplicate_count": sum(1 for rows in raw_failure_rows.values() if len(rows) > 1),
        "raw_failure_id_cross_family_conflicts": {
            key: sorted(value) for key, value in raw_failure.items() if len(value) > 1
        },
        "raw_hazard_event_id_reuse_count": sum(1 for values in raw_hazard.values() if len(values) > 1),
        "raw_safety_goal_id_reuse_count": sum(1 for values in raw_sg.values() if len(values) > 1),
        "semantic_key_duplicate_count": len(records) - len({record["semantic_key"] for record in records}),
        "unresolved_function_rows": [record["trace"]["source_row"] for record in records if record["function_match_status"] == "unknown"],
        "unresolved_hazard_rows": [record["trace"]["source_row"] for record in records if record["hazard_mapping_status"] == "unresolved"],
        "invalid_sec_rows": [record["trace"]["source_row"] for record in records if record["asil_status"] == "invalid_sec_combination"],
    }



def _vehicle_safety_goal_catalog(workbook, input_path: Path) -> dict[str, Any] | None:
    """Read the optional flat ``整车安全目标`` reference sheet.

    The case table remains the source of event semantics; this catalog only
    preserves the reviewed vehicle-level SG IDs and their display metadata so
    S5 can use them as an explicit prefill mapping.
    """
    if "整车安全目标" not in workbook.sheetnames:
        return None
    ws = workbook["整车安全目标"]
    headers: dict[str, int] = {}
    aliases = {
        "id": {"整车安全目标id"},
        "safety_goal": {"整车安全目标"},
        "safe_state": {"安全状态"},
        "ftti": {"ftti"},
        "note": {"注释", "备注"},
        "filled_by": {"填写人"},
        "filled_at": {"填写时间"},
    }
    for col in range(1, ws.max_column + 1):
        key = header_key(ws.cell(1, col).value)
        for logical, candidates in aliases.items():
            if key in candidates and logical not in headers:
                headers[logical] = col
    if "id" not in headers or "safety_goal" not in headers:
        raise ValueError("整车安全目标 Sheet 缺少整车安全目标ID或整车安全目标列")

    goals: dict[str, dict[str, Any]] = {}
    duplicate_ids: list[str] = []
    for row_number in range(2, ws.max_row + 1):
        goal_id = normalize_text(ws.cell(row_number, headers["id"]).value)
        if not goal_id:
            continue
        item = {
            "vehicle_safety_goal_id": goal_id,
            "safety_goal": normalize_text(ws.cell(row_number, headers["safety_goal"]).value) or None,
            "safe_state": normalize_text(ws.cell(row_number, headers["safe_state"]).value) if "safe_state" in headers else None,
            "ftti": normalize_text(ws.cell(row_number, headers["ftti"]).value) if "ftti" in headers else None,
            "note": normalize_text(ws.cell(row_number, headers["note"]).value) if "note" in headers else None,
            "source_ref": {
                "source_file": input_path.name,
                "source_sheet": ws.title,
                "source_row": row_number,
            },
        }
        if goal_id in goals:
            duplicate_ids.append(goal_id)
        goals[goal_id] = item
    if duplicate_ids:
        raise ValueError(f"整车安全目标 Sheet 存在重复 ID: {sorted(set(duplicate_ids))}")
    return {
        "schema_version": "pt_vehicle_safety_goal_reference.v1",
        "source": {
            "file": str(input_path),
            "sheet": ws.title,
            "row_count": len(goals),
        },
        "goals": dict(sorted(goals.items())),
    }

def import_case_table(input_path: str | Path, *, sheet_name: str | None = None, hazard_catalog: dict[str, Any] | None = None) -> tuple[dict[str, Any], dict[str, Any]]:
    from openpyxl import load_workbook

    input_path = Path(input_path).resolve()
    workbook = load_workbook(input_path, data_only=True, read_only=False)
    if sheet_name:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"找不到 Sheet: {sheet_name}")
        ws = workbook[sheet_name]
    elif "案例表" in workbook.sheetnames:
        ws = workbook["案例表"]
    elif len(workbook.worksheets) == 1:
        ws = workbook.active
    else:
        raise ValueError(f"无法确定案例表 Sheet，可选: {workbook.sheetnames}")

    headers = _make_headers(ws)
    vehicle_goal_catalog = _vehicle_safety_goal_catalog(workbook, input_path)
    required = {"function_name", "failure_mode", "anomaly", "vehicle_hazard", "scenario", "severity", "exposure", "controllability"}
    missing = sorted(required - set(headers))
    if missing:
        raise ValueError(f"案例表缺少字段: {missing}")

    if hazard_catalog is None:
        try:
            sys.path.insert(0, str(_SCRIPT_DIR))
            from utils.domain_packs import load_domain_pack
            pack = load_domain_pack("PT")
            hazard_catalog = pack.get("safety_goal_catalog", {}).get("hazard_catalog", {})
        except Exception:
            pack = None
            hazard_catalog = {}
    else:
        pack = None

    records: list[dict[str, Any]] = []
    for row_number in range(2, ws.max_row + 1):
        raw = _raw_row(ws, headers, row_number)
        if not any(normalize_text(value) for value in raw.values()):
            continue
        records.append(
            normalize_row(
                raw,
                source_file=input_path.name,
                source_sheet=ws.title,
                source_row=row_number,
                hazard_catalog=hazard_catalog,
                pt_pack=pack if "pack" in locals() else None,
            )
        )

    if vehicle_goal_catalog is not None:
        goal_ids = set(vehicle_goal_catalog["goals"])
        referenced_ids = {
            record.get("vehicle_safety_goal_id")
            for record in records
            if record.get("vehicle_safety_goal_id")
        }
        unresolved = sorted(referenced_ids - goal_ids)
        if unresolved:
            raise ValueError(
                "案例表引用了未在整车安全目标 Sheet 定义的 ID: "
                + ", ".join(unresolved)
            )

    audit = audit_records(records)
    inventory = {
        "schema_version": "pt_reference_case_source.v1",
        "source_policy": {
            "source_ids_are_trace_only": True,
            "runtime_matching_keys": ["canonical_function_family", "semantic_role", "normalized_failure_mode", "anomaly", "vehicle_hazard_id", "scenario_id"],
            "asil_is_recomputed": True,
        },
        "source": {
            "file": str(input_path),
            "sheet": ws.title,
            "row_count": len(records),
        },
        "records": records,
        "vehicle_safety_goal_catalog": vehicle_goal_catalog,
        "audit_summary": audit,
    }
    return inventory, audit
