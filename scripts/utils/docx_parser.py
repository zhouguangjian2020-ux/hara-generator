# utils/docx_parser.py
# 文档解析：功能清单快路径 + H1～H8 最小上下文单遍提取
# 无 CLI，无外部副作用

from __future__ import annotations

from pathlib import Path
import re
from typing import Any

from docx import Document

from utils.data_models import RelatedItem, SubFunction, extract_domain
from utils.document_context import (
    DOCUMENT_CONTEXT_SCHEMA,
    MAX_NODE_BODY_CHARS,
    MAX_NODE_BODY_PARAGRAPHS,
    build_node_summary_map,
    classify_section,
)


_HEADER_ALIASES = {
    "func_id": ("相关项ID", "整车功能ID"),
    "func_name": ("相关项", "整车层级功能"),
    "sub_name": ("子功能",),
    "scenario": ("相关项中的子功能",),
}
_REQUIRED_COLUMNS = ("func_id", "func_name", "feature_id", "sub_name")
_DESCRIPTION_MARKERS = {"功能描述", "功能详细定义", "功能需求"}
_FUNCTION_ROOT_MARKERS = ("功能描述", "整车功能定义", "功能定义")
_SKIP_BODY_STYLES = {"Caption", "标题", "页眉", "页脚"}
_LEADING_NUMBER_RE = re.compile(r"^[\s.．、\-—_0-9一二三四五六七八九十]+")


class FunctionTableNotFoundError(ValueError):
    """未找到可识别功能清单时抛出；携带候选表头供 Agent 回退。"""

    def __init__(self, message: str, *, table_headers: list[dict] | None = None):
        super().__init__(message)
        self.table_headers = table_headers or []


def parse_docx(path: Path, domain: str = "") -> tuple[str, list[RelatedItem], list[dict]]:
    """兼容旧接口：返回 (文档名, 相关项, legacy章节)。"""
    doc_name, related_items, chapters, _context, _parser = parse_docx_with_context(
        path, domain=domain,
    )
    return doc_name, related_items, chapters


def parse_docx_with_context(
    path: Path,
    domain: str = "",
    *,
    source_path: Path | None = None,
) -> tuple[str, list[RelatedItem], list[dict], dict, str]:
    """解析 DOCX，并返回增量 document_context_v1。"""
    parse_path = Path(path)
    logical_source = Path(source_path) if source_path else parse_path
    doc = Document(str(parse_path))
    doc_name = _extract_doc_name(logical_source)

    standard = _find_standard_function_table(doc)
    if standard is not None:
        table_index, table = standard
        related_items = _parse_function_table(
            table,
            explicit_domain=domain,
            source_path=logical_source,
            source_table_index=table_index,
        )
        parser_name = "standard_feature_list"
        table_evidence = [{
            "kind": "function_list",
            "table_index": table_index,
            "headers": [cell.text.strip() for cell in table.rows[0].cells],
        }]
    else:
        generic = _find_generic_function_summary_table(doc)
        if generic is None:
            context = _extract_document_context(doc)
            raise FunctionTableNotFoundError(
                "未找到可识别功能清单表；需要 Agent 结构化解析",
                table_headers=_table_header_samples(doc),
            )
        summary_index, summary_table = generic
        related_items, table_evidence = _parse_generic_function_tables(
            doc,
            summary_index,
            summary_table,
            explicit_domain=domain,
            source_path=logical_source,
        )
        parser_name = "generic_function_definition"

    chapters, document_context = _parse_chapter3_with_context(doc)
    document_context["table_evidence"] = table_evidence
    document_context["parser"] = parser_name
    return doc_name, related_items, chapters, document_context, parser_name


def inspect_docx_context(path: Path, *, source_path: Path | None = None) -> tuple[str, dict, list[dict]]:
    """即使功能表无法识别，也提取受限目录和正文供 Agent 回退。"""
    parse_path = Path(path)
    logical_source = Path(source_path) if source_path else parse_path
    doc = Document(str(parse_path))
    _chapters, context = _parse_chapter3_with_context(doc)
    context["table_evidence"] = _table_header_samples(doc)
    return _extract_doc_name(logical_source), context, _table_header_samples(doc)


# ========== 文档名 ==========

def _extract_doc_name(path: Path) -> str:
    """返回输入文档的原始文件名，不做任何清洗或重命名。

    该值会写入S1“相关项功能清单”的“相关项定义文档”列，
    因此必须保留完整文件名及扩展名。
    """
    return path.name


# ========== 功能清单 ==========

def _table_header_samples(doc, limit: int = 20) -> list[dict]:
    samples = []
    for table_index, table in enumerate(doc.tables[:limit]):
        if not table.rows:
            continue
        samples.append({
            "table_index": table_index,
            "row_count": len(table.rows),
            "column_count": len(table.columns),
            "headers": [cell.text.strip()[:120] for cell in table.rows[0].cells],
        })
    return samples


def _find_standard_function_table(doc):
    for table_index, table in enumerate(doc.tables):
        if not table.rows:
            continue
        header_cells = [cell.text.strip() for cell in table.rows[0].cells]
        has_item = any("相关项" in header for header in header_cells)
        has_feature = any("子功能" in header and "ID" in header for header in header_cells)
        if has_item and has_feature:
            return table_index, table
    return None


def _find_function_table(doc):
    """保留旧的内部接口；仅返回标准模板功能表。"""
    found = _find_standard_function_table(doc)
    if found is None:
        raise FunctionTableNotFoundError(
            "未找到相关项功能清单表（需包含'相关项'与'子功能ID'表头）",
            table_headers=_table_header_samples(doc),
        )
    return found[1]


def _find_generic_function_summary_table(doc):
    for table_index, table in enumerate(doc.tables):
        if not table.rows:
            continue
        headers = [cell.text.strip().replace(" ", "") for cell in table.rows[0].cells]
        if "功能ID" in headers and "功能名称" in headers:
            return table_index, table
    return None


def _locate_columns(headers: list[str]) -> dict:
    cols = {}
    for key, aliases in _HEADER_ALIASES.items():
        cols[key] = next((index for index, header in enumerate(headers) if header in aliases), None)
    cols["feature_id"] = next(
        (index for index, header in enumerate(headers) if "子功能" in header and "ID" in header),
        None,
    )
    if cols["func_id"] is None:
        cols["func_id"] = next(
            (index for index, header in enumerate(headers) if "相关项" in header and "ID" in header),
            None,
        )
    return cols


def _append_component(
    sub: SubFunction,
    name: str,
    scenario: str,
    source_row: int,
    *,
    source_table_index: int | None = None,
) -> None:
    name = (name or "").strip()
    scenario = (scenario or "").strip()
    if not name and not scenario:
        return
    component = {
        "name": name,
        "scenario": scenario,
        "source_row": source_row,
    }
    if source_table_index is not None:
        component["source_table_index"] = source_table_index
    key = (name, scenario)
    if not any(
        (candidate.get("name", ""), candidate.get("scenario", "")) == key
        for candidate in sub.components
    ):
        sub.components.append(component)


def _parse_function_table(
    table,
    *,
    explicit_domain: str = "",
    source_path: Path | None = None,
    source_table_index: int | None = None,
) -> list[RelatedItem]:
    headers = [cell.text.strip() for cell in table.rows[0].cells]
    cols = _locate_columns(headers)
    missing = [key for key in _REQUIRED_COLUMNS if cols.get(key) is None]
    if missing:
        raise ValueError(f"功能清单表缺少必需列 {list(missing)}，实际表头为: {headers}")

    items: list[RelatedItem] = []
    current: RelatedItem | None = None
    feature_map: dict[str, SubFunction] = {}
    for row_index, row in enumerate(table.rows[1:], start=2):
        cells = [cell.text.strip() for cell in row.cells]
        func_id = cells[cols["func_id"]].replace(" ", "")
        func_name = cells[cols["func_name"]]
        feature_id = cells[cols["feature_id"]]
        sub_name = cells[cols["sub_name"]]
        scenario = cells[cols["scenario"]] if cols.get("scenario") is not None else ""
        if not func_id and not feature_id:
            continue
        if func_id and (current is None or current.func_id != func_id):
            current = RelatedItem(func_id=func_id, func_name=func_name)
            items.append(current)
            feature_map = {}
        if current is None:
            continue
        feature_key = feature_id or sub_name or scenario
        if not feature_key:
            continue
        sub = feature_map.get(feature_key)
        if sub is None:
            generated_id = feature_id or f"{current.func_id}_S{len(current.sub_functions) + 1:02d}"
            sub = SubFunction(
                feature_list_id=generated_id,
                name=sub_name or scenario,
                description="",
                chapter="",
                scenario=scenario,
                source_table_index=source_table_index,
                source_row=row_index,
                evidence_refs=[f"table:{source_table_index}:row:{row_index}"],
            )
            feature_map[feature_key] = sub
            current.sub_functions.append(sub)
        _append_component(
            sub,
            sub_name,
            scenario,
            row_index,
            source_table_index=source_table_index,
        )
        if scenario and scenario not in sub.scenario:
            sub.scenario = f"{sub.scenario}；{scenario}" if sub.scenario else scenario

    _resolve_item_domains(items, explicit_domain, source_path)
    return items


def _cell_map(table) -> dict[str, str]:
    values: dict[str, str] = {}
    for row in table.rows:
        cells = [cell.text.strip() for cell in row.cells]
        if len(cells) >= 2 and cells[0]:
            values[cells[0].replace(" ", "")] = cells[1]
    return values


def _bounded(value: str, max_chars: int = 1800) -> str:
    value = str(value or "").strip()
    return value if len(value) <= max_chars else value[:max_chars] + "…"


def _parse_generic_function_tables(
    doc,
    summary_index: int,
    summary_table,
    *,
    explicit_domain: str,
    source_path: Path | None,
) -> tuple[list[RelatedItem], list[dict]]:
    headers = [cell.text.strip().replace(" ", "") for cell in summary_table.rows[0].cells]
    id_col = headers.index("功能ID")
    name_col = headers.index("功能名称")
    overview_col = headers.index("功能概述") if "功能概述" in headers else None
    items: list[RelatedItem] = []
    by_id: dict[str, RelatedItem] = {}
    overviews: dict[str, str] = {}
    evidence = [{
        "kind": "generic_function_summary",
        "table_index": summary_index,
        "headers": headers,
    }]
    for row_index, row in enumerate(summary_table.rows[1:], start=2):
        cells = [cell.text.strip() for cell in row.cells]
        func_id = cells[id_col].replace(" ", "")
        func_name = cells[name_col]
        if not func_id and not func_name:
            continue
        item = RelatedItem(func_id=func_id, func_name=func_name)
        items.append(item)
        by_id[func_id] = item
        overviews[func_id] = cells[overview_col] if overview_col is not None else ""

    for table_index, table in enumerate(doc.tables):
        if table_index == summary_index or len(table.rows) < 2:
            continue
        values = _cell_map(table)
        feature_id = str(values.get("功能ID", "")).strip().replace(" ", "")
        feature_name = str(values.get("功能", "")).strip()
        if not feature_id or not feature_name:
            continue
        owner_id = feature_id.split("_", 1)[0]
        owner = by_id.get(owner_id)
        if owner is None and len(items) == 1:
            owner = items[0]
        if owner is None:
            continue
        description = _bounded(values.get("功能需求", "") or values.get("功能详细定义", ""))
        scenario = str(values.get("运行模式", "") or values.get("工作模式", "")).strip()
        sub = SubFunction(
            feature_list_id=feature_id,
            name=feature_name,
            description=description,
            chapter="",
            scenario=scenario,
            source_table_index=table_index,
            source_row=1,
            evidence_refs=[
                f"table:{table_index}:row:1",
                f"table:{table_index}:row:2",
                f"table:{table_index}:row:3",
            ],
        )
        _append_component(
            sub,
            feature_name,
            scenario,
            2,
            source_table_index=table_index,
        )
        owner.sub_functions.append(sub)
        evidence.append({
            "kind": "generic_function_detail",
            "table_index": table_index,
            "feature_list_id": feature_id,
            "headers": [row.cells[0].text.strip() for row in table.rows[:4]],
        })

    # 没有明细表时仍允许按整车功能名称进入 S1 快路径；ID 由来源功能ID复用。
    for item in items:
        if not item.sub_functions:
            item.sub_functions.append(SubFunction(
                feature_list_id=item.func_id or "SOURCE_FEATURE_0001",
                name=item.func_name,
                description=_bounded(overviews.get(item.func_id, "")),
                chapter="",
                source_table_index=summary_index,
                source_row=2,
                evidence_refs=[f"table:{summary_index}:row:2"],
            ))
    _resolve_item_domains(items, explicit_domain, source_path)
    return items, evidence


def _resolve_item_domains(items: list[RelatedItem], explicit_domain: str, source_path: Path | None) -> None:
    for item in items:
        semantic_texts = [item.func_name]
        for sub in item.sub_functions:
            semantic_texts.extend([sub.name, sub.description, sub.scenario])
            for component in sub.components:
                semantic_texts.extend([component.get("name", ""), component.get("scenario", "")])
        item.domain = extract_domain(
            item.func_id,
            explicit_domain=explicit_domain,
            semantic_texts=semantic_texts,
            source_path=source_path,
        )


# ========== 单遍章节与上下文提取 ==========

def _paragraph_style_name(paragraph, style_names: dict[str, str]) -> str:
    """直接读取 pStyle ID，避免对每个段落调用 python-docx 的昂贵样式查找。"""
    properties = paragraph._p.pPr
    if properties is None or properties.pStyle is None or properties.pStyle.val is None:
        return "Normal"
    style_id = str(properties.pStyle.val)
    return style_names.get(style_id, style_id)


def _heading_level(paragraph, style_name: str | None = None) -> int | None:
    """返回标题层级 1～8；兼容 Heading 样式和 outlineLvl。"""
    name = style_name if style_name is not None else (paragraph.style.name or "")
    if name.startswith("Heading"):
        try:
            level = int(name.split()[-1])
            return level if 1 <= level <= 8 else None
        except ValueError:
            pass
    paragraph_properties = paragraph._p.pPr
    if paragraph_properties is not None:
        outline = paragraph_properties.outlineLvl
        if outline is not None and outline.val is not None:
            level = int(outline.val) + 1
            return level if 1 <= level <= 8 else None
    return None


def _clean_heading(text: str) -> str:
    return _LEADING_NUMBER_RE.sub("", str(text or "").strip())


def _is_function_root(text: str) -> bool:
    cleaned = _clean_heading(text)
    return any(marker in cleaned for marker in _FUNCTION_ROOT_MARKERS)


def _should_skip_body(style: str) -> bool:
    return style in _SKIP_BODY_STYLES or style.lower().startswith("toc")


def _nearest_parent(stack: dict[int, str], level: int) -> str | None:
    lower = [candidate for candidate in stack if candidate < level]
    return stack[max(lower)] if lower else None


def _parse_chapter3(doc: Document) -> list[dict]:
    chapters, _context = _parse_chapter3_with_context(doc)
    return chapters


def _parse_chapter3_with_context(doc: Document) -> tuple[list[dict], dict]:
    chapters: list[dict] = []
    nodes: list[dict] = []
    node_stack: dict[int, str] = {}
    heading_text_stack: dict[int, str] = {}
    numbering = [0] * 9
    current_node: dict | None = None

    h1_num = 0
    function_chapter_num = 0
    in_function_root = False
    h2_num = h3_num = h4_num = 0
    current_h2 = ""
    current_h3: dict | None = None
    current_h4: dict | None = None

    style_names = {style.style_id: style.name for style in doc.styles}
    for source_index, paragraph in enumerate(doc.paragraphs):
        style = _paragraph_style_name(paragraph, style_names)
        level = _heading_level(paragraph, style)
        text = paragraph.text.strip()
        if level is not None:
            if level == 1:
                h1_num += 1
                in_function_root = _is_function_root(text)
                if in_function_root:
                    function_chapter_num = h1_num
                    h2_num = h3_num = h4_num = 0
                    current_h2 = ""
                    current_h3 = current_h4 = None
            for deeper in range(level + 1, 9):
                numbering[deeper] = 0
                node_stack.pop(deeper, None)
                heading_text_stack.pop(deeper, None)
            numbering[level] += 1
            heading_text_stack[level] = text
            heading_path = [
                heading_text_stack[candidate]
                for candidate in sorted(heading_text_stack)
                if candidate <= level
            ]
            node_id = f"n{len(nodes) + 1:04d}"
            label_numbers = [str(numbering[candidate]) for candidate in range(1, level + 1) if numbering[candidate]]
            label = f"{'.'.join(label_numbers)} {text}".strip()
            node = {
                "node_id": node_id,
                "parent_id": _nearest_parent(node_stack, level),
                "source_index": source_index,
                "level": level,
                "label": label,
                "heading_text": text,
                "section_type": classify_section(heading_path),
                "body": [],
                "evidence_ref": f"paragraph:{source_index}",
            }
            nodes.append(node)
            node_stack[level] = node_id
            current_node = node

            if not in_function_root or level == 1:
                continue
            if level == 2:
                h2_num += 1
                h3_num = h4_num = 0
                current_h2 = text
                current_h3 = current_h4 = None
            elif level == 3:
                h3_num += 1
                h4_num = 0
                if _clean_heading(text) in {"功能概述", "功能概要"}:
                    current_h3 = current_h4 = None
                else:
                    legacy_label = f"{function_chapter_num}.{h2_num}.{h3_num} {text}"
                    current_h3 = {
                        "h2": current_h2,
                        "h3": text,
                        "label": legacy_label,
                        "desc": None,
                        "h4s": [],
                        "_node_id": node_id,
                    }
                    chapters.append(current_h3)
                    current_h4 = None
            elif level == 4:
                h4_num += 1
                if current_h3 is None or _clean_heading(text) in _DESCRIPTION_MARKERS:
                    current_h4 = None
                else:
                    legacy_label = f"{function_chapter_num}.{h2_num}.{h3_num}.{h4_num} {text}"
                    current_h4 = {
                        "h4": text,
                        "label": legacy_label,
                        "desc": None,
                        "_node_id": node_id,
                    }
                    current_h3["h4s"].append(current_h4)
            continue

        if not text or current_node is None or _should_skip_body(style):
            continue
        if len(current_node["body"]) >= MAX_NODE_BODY_PARAGRAPHS:
            continue
        used_chars = sum(len(item.get("text", "")) for item in current_node["body"])
        remain = MAX_NODE_BODY_CHARS - used_chars
        if remain <= 0:
            continue
        current_node["body"].append({
            "style": style,
            "text": text[:remain],
            "source_index": source_index,
            "evidence_ref": f"paragraph:{source_index}",
        })

    document_context = {
        "schema_version": DOCUMENT_CONTEXT_SCHEMA,
        "nodes": nodes,
        "metrics": {
            "heading_node_count": len(nodes),
            "body_node_count": sum(1 for node in nodes if node.get("body")),
            "body_char_count": sum(
                len(item.get("text", ""))
                for node in nodes
                for item in node.get("body", [])
            ),
        },
    }

    summary_map = build_node_summary_map(document_context, max_chars=1200)
    for chapter in chapters:
        chapter["desc"] = summary_map.get(chapter.pop("_node_id"), "") or None
        for h4 in chapter.get("h4s", []):
            h4["desc"] = summary_map.get(h4.pop("_node_id"), "") or None
    return chapters, document_context
