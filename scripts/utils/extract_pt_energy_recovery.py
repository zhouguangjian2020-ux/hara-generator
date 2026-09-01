#!/usr/bin/env python3
"""
从山子高科PT域Excel提取能量回收功能数据
"""

import openpyxl
import sys
import json
from pathlib import Path

sys.stdout.reconfigure(encoding='utf-8')


def get_cell_value(ws, row, col, merged_ranges):
    """获取单元格值，考虑合并单元格"""
    cell = ws.cell(row=row, column=col)
    value = cell.value

    # 如果当前单元格在合并区域内，获取合并区域的起始单元格值
    if value is None:
        for merged in merged_ranges:
            if merged.min_row <= row <= merged.max_row and merged.min_col <= col <= merged.max_col:
                value = ws.cell(row=merged.min_row, column=merged.min_col).value
                break

    return value


def extract_hara_data(excel_path):
    """提取HARA分析数据"""
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws = wb['HARA分析  ']

    merged_ranges = list(ws.merged_cells.ranges)

    # 提取数据
    events = []

    # 从第4行开始（第3行是功能标题）
    for r in range(4, ws.max_row + 1):
        a_val = ws.cell(row=r, column=1).value
        if not a_val:
            continue

        event = {
            'hazard_id': get_cell_value(ws, r, 1, merged_ranges),
            'sub_function': get_cell_value(ws, r, 2, merged_ranges),
            'failure_id': get_cell_value(ws, r, 3, merged_ranges),
            'anomaly': get_cell_value(ws, r, 4, merged_ranges),
            'vehicle_hazard': get_cell_value(ws, r, 5, merged_ranges),
            'scenario': get_cell_value(ws, r, 6, merged_ranges),
            'description': get_cell_value(ws, r, 7, merged_ranges),
            'severity': get_cell_value(ws, r, 8, merged_ranges),
            'severity_reason': get_cell_value(ws, r, 9, merged_ranges),
            'exposure': get_cell_value(ws, r, 10, merged_ranges),
            'exposure_reason': get_cell_value(ws, r, 11, merged_ranges),
            'controllability': get_cell_value(ws, r, 12, merged_ranges),
            'controllability_reason': get_cell_value(ws, r, 13, merged_ranges),
            'asil': get_cell_value(ws, r, 14, merged_ranges),
            'safety_goal_id': get_cell_value(ws, r, 15, merged_ranges),
            'safety_goal': get_cell_value(ws, r, 16, merged_ranges),
            'safe_state': get_cell_value(ws, r, 17, merged_ranges),
            'ftti': get_cell_value(ws, r, 18, merged_ranges),
            'note': get_cell_value(ws, r, 19, merged_ranges)
        }
        events.append(event)

    return events


def extract_hazop_data(excel_path):
    """提取HAZOP分析数据"""
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws = wb['HAZOP 分析']

    merged_ranges = list(ws.merged_cells.ranges)

    hazop_entries = []

    # 从第3行开始（第2行是表头）
    for r in range(3, ws.max_row + 1):
        # HAZOP表使用了大量合并单元格，需要用get_cell_value处理
        func_id = get_cell_value(ws, r, 1, merged_ranges)  # A - 功能ID
        func_name = get_cell_value(ws, r, 2, merged_ranges)  # B - 功能名
        failure_mode = get_cell_value(ws, r, 3, merged_ranges)  # C - 失效模式
        anomaly = get_cell_value(ws, r, 4, merged_ranges)  # D - 功能异常表现
        failure_id = get_cell_value(ws, r, 5, merged_ranges)  # E - 功能失效ID
        vehicle_hazard = get_cell_value(ws, r, 6, merged_ranges)  # F - 整车危害
        related_hara = get_cell_value(ws, r, 7, merged_ranges)  # G - 关联HARA

        # E列有值的才是有效数据行
        if not failure_id:
            continue

        entry = {
            'function_id': func_id,
            'function_name': func_name,
            'failure_mode': failure_mode,
            'anomaly': anomaly,
            'failure_id': failure_id,
            'vehicle_hazard': vehicle_hazard,
            'related_hara': related_hara
        }
        hazop_entries.append(entry)

    return hazop_entries


def extract_safety_goals(excel_path):
    """提取整车安全目标"""
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws = wb['整车安全目标']

    merged_ranges = list(ws.merged_cells.ranges)

    safety_goals = []

    # 从第5行开始（前4行是表头）
    for r in range(5, ws.max_row + 1):
        a_val = ws.cell(row=r, column=1).value
        if not a_val:
            continue

        sg = {
            'event_sg_id': get_cell_value(ws, r, 1, merged_ranges),
            'safety_goal': get_cell_value(ws, r, 2, merged_ranges),
            'asil': get_cell_value(ws, r, 3, merged_ranges),
            'safe_state': get_cell_value(ws, r, 4, merged_ranges),
            'ftti': get_cell_value(ws, r, 5, merged_ranges),
            'vehicle_sg_id': get_cell_value(ws, r, 7, merged_ranges),
            'vehicle_safety_goal': get_cell_value(ws, r, 8, merged_ranges),
            'vehicle_asil': get_cell_value(ws, r, 9, merged_ranges),
            'vehicle_safe_state': get_cell_value(ws, r, 10, merged_ranges),
            'vehicle_ftti': get_cell_value(ws, r, 11, merged_ranges)
        }
        safety_goals.append(sg)

    return safety_goals


def analyze_data(hara_events, hazop_entries, safety_goals):
    """分析提取的数据"""
    from collections import Counter

    print('=== 能量回收功能数据提取汇总 ===\n')

    # 1. HARA事件统计
    print(f'【HARA事件】总数: {len(hara_events)}\n')

    # 按子功能分组
    sub_funcs = Counter(e['sub_function'] for e in hara_events if e['sub_function'])
    print('按子功能分布:')
    for sf, count in sub_funcs.most_common():
        print(f'  {sf}: {count}个事件')

    # 按失效模式分组
    failure_ids = Counter(e['failure_id'] for e in hara_events if e['failure_id'])
    print(f'\n失效ID数量: {len(failure_ids)}')

    # ASIL分布
    asil_dist = Counter(e['asil'] for e in hara_events)
    print('\nASIL分布:')
    for asil in ['QM', 'A', 'B', 'C', 'D', None]:
        count = asil_dist.get(asil, 0)
        if count > 0:
            print(f'  {asil if asil else "空"}: {count}')

    # 2. HAZOP条目统计
    print(f'\n【HAZOP分析】总数: {len(hazop_entries)}\n')

    # 失效模式
    failure_modes = Counter(e['failure_mode'] for e in hazop_entries if e['failure_mode'])
    print('失效模式分布:')
    for fm, count in failure_modes.most_common():
        print(f'  {fm}: {count}条')

    # 功能异常表现
    anomalies = Counter(e['anomaly'] for e in hazop_entries if e['anomaly'])
    print('\n功能异常表现分布:')
    for anomaly, count in list(anomalies.most_common())[:5]:
        print(f'  {anomaly}: {count}条')

    # 3. 安全目标统计
    print(f'\n【整车安全目标】总数: {len(safety_goals)}\n')

    # 去重整车级SG
    vehicle_sgs = set((sg['vehicle_sg_id'], sg['vehicle_safety_goal'])
                      for sg in safety_goals if sg['vehicle_sg_id'])
    print(f'唯一整车级安全目标: {len(vehicle_sgs)}个')

    # 4. 数据质量检查
    print('\n=== 数据质量检查 ===\n')

    # 检查S/E/C完整性
    complete = sum(1 for e in hara_events
                   if e['severity'] is not None
                   and e['exposure'] is not None
                   and e['controllability'] is not None)
    print(f'S/E/C完整的事件: {complete}/{len(hara_events)} ({complete/len(hara_events)*100:.1f}%)')

    # 检查安全目标
    with_sg = sum(1 for e in hara_events if e['safety_goal'])
    print(f'包含安全目标的事件: {with_sg}/{len(hara_events)} ({with_sg/len(hara_events)*100:.1f}%)')

    # 检查场景描述
    with_scenario = sum(1 for e in hara_events if e['scenario'])
    print(f'包含场景描述的事件: {with_scenario}/{len(hara_events)} ({with_scenario/len(hara_events)*100:.1f}%)')


def save_json(data, output_path):
    """保存为JSON"""
    with open(output_path, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    print(f'\n已保存: {output_path}')


if __name__ == '__main__':
    excel_path = 'C:/Users/zhouguangjian/Downloads/FUSA HARA PT DOMAIN_V1.1_山子高科 - 副本.xlsx'

    # 提取数据
    print('正在提取数据...\n')
    hara_events = extract_hara_data(excel_path)
    hazop_entries = extract_hazop_data(excel_path)
    safety_goals = extract_safety_goals(excel_path)

    # 分析数据
    analyze_data(hara_events, hazop_entries, safety_goals)

    # 保存为JSON
    output = {
        'function_id': 'P_func_0008',  # 新功能ID
        'function_name': '能量回收功能',
        'source': 'FUSA HARA PT DOMAIN_V1.1_山子高科',
        'hara_events': hara_events,
        'hazop_entries': hazop_entries,
        'safety_goals': safety_goals
    }

    save_json(output, 'output/pt_energy_recovery_extracted.json')
