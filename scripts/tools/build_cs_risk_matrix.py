"""从一份已批准的 CS 转向 HARA 参考中提炼语义风险矩阵到 CS Domain Pack。

这是离线知识沉淀工具：只读取 HARA 分析 Sheet 的业务字段，不依赖参考文件的
列号、封面、其他 Sheet 或样式。运行时不会读取参考 Excel。
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
DEFAULT_PACK = ROOT / "references" / "domain_packs" / "CS.json"


def _norm(value) -> str:
    return str(value or "").replace("\n", "").replace(" ", "").strip()


def _headers(ws) -> dict[str, int]:
    result: dict[str, int] = {}
    for column in range(1, ws.max_column + 1):
        for row in (1, 2):
            value = _norm(ws.cell(row, column).value)
            if value:
                result.setdefault(value, column)
    return result


def _column(headers: dict[str, int], exact: str, *contains: str) -> int:
    if exact in headers:
        return headers[exact]
    for header, column in headers.items():
        if all(token in header for token in contains):
            return column
    raise ValueError(f"HARA Sheet 缺少逻辑字段：{exact}")


def extract(reference: Path) -> tuple[list[dict], list[dict]]:
    wb = load_workbook(reference, data_only=False, read_only=True)
    if "HARA 分析" not in wb.sheetnames:
        raise ValueError("参考文件缺少 HARA 分析 Sheet")
    ws = wb["HARA 分析"]
    headers = _headers(ws)
    columns = {
        "failure_id": _column(headers, "功能失效ID", "功能失效", "ID"),
        "scenario": _column(headers, "运行场景(驾驶和运行情况)", "运行场景"),
        "description": _column(headers, "危害事件描述", "危害事件描述"),
        "severity": _column(headers, "严重度(S)", "严重度"),
        "severity_reason": _column(headers, "对“S”的理由", "对", "S", "理由"),
        "exposure": _column(headers, "暴露概率(E)", "暴露概率"),
        "exposure_reason": _column(headers, "对E的理由", "对", "E", "理由"),
        "controllability": _column(headers, "可控性（C）", "可控性"),
        "controllability_reason": _column(headers, "对C的理由", "对", "C", "理由"),
        "asil": _column(headers, "ASIL", "ASIL"),
        "safety_goal": _column(headers, "安全目标", "安全目标"),
        "safe_state": _column(headers, "安全状态", "安全状态"),
        "ftti": _column(headers, "FTTI", "FTTI"),
    }

    rows: list[dict] = []
    scenario_order: list[str] = []
    for row in range(3, ws.max_row + 1):
        failure_id = ws.cell(row, columns["failure_id"]).value
        scenario = ws.cell(row, columns["scenario"]).value
        if not isinstance(failure_id, str) or not failure_id.startswith("CS_MF_"):
            continue
        if not isinstance(scenario, str) or not scenario.strip():
            continue
        scenario = scenario.strip()
        if scenario not in scenario_order:
            scenario_order.append(scenario)
        item = {key: ws.cell(row, column).value for key, column in columns.items()}
        item["failure_id"] = failure_id.strip()
        item["scenario"] = scenario
        rows.append(item)

    if len(scenario_order) != 19:
        raise ValueError(f"CS 转向基线应为 19 个场景，实际为 {len(scenario_order)}")
    if len(rows) != 95:
        raise ValueError(f"CS 转向基线应为 95 条事件，实际为 {len(rows)}")
    if {item["failure_id"] for item in rows} != {
        "CS_MF_0001_01", "CS_MF_0001_02", "CS_MF_0001_03", "CS_MF_0001_04", "CS_MF_0001_05"
    }:
        raise ValueError("CS 转向基线 failure_id 集合不符合五模式标准")

    scenarios = [
        {"scenario_id": f"CS_SC_{index:02d}", "text": scenario}
        for index, scenario in enumerate(scenario_order, start=1)
    ]
    id_by_text = {item["text"]: item["scenario_id"] for item in scenarios}
    return scenarios, [
        {
            "analysis_unit_id": "CS_STEERING_ASSIST_MAIN",
            "failure_id": item["failure_id"],
            "scenario_id": id_by_text[item["scenario"]],
            "description": item["description"],
            "severity": item["severity"],
            "severity_reason": item["severity_reason"],
            "exposure": item["exposure"],
            "exposure_reason": item["exposure_reason"],
            "controllability": item["controllability"],
            "controllability_reason": item["controllability_reason"],
            # 项目质量合同规定 S=0 时 E/C 必须为空，ASIL 统一留空，
            # 因此不能直接沿用历史参考中常见的 QM 单元格文本。
            "expected_asil": None if item["severity"] == 0 else item["asil"],
            "safety_goal": item["safety_goal"],
            "safe_state": item["safe_state"],
            "ftti": item["ftti"],
        }
        for item in rows
    ]


def update_pack(reference: Path, pack_path: Path, dry_run: bool = False) -> None:
    scenarios, matrix = extract(reference)
    pack = json.loads(pack_path.read_text(encoding="utf-8"))

    # 整车危害文本不从参考 Excel 继承：它是 Pack 内的稳定业务资产。
    # 每条风险事件只通过 failure_id 继承对应的 vehicle_hazard_id。
    hazard_by_failure = {
        pattern.get("failure_id"): pattern.get("vehicle_hazard_id")
        for pattern in pack.get("analysis_catalog", {}).get("hazop_patterns", [])
        if isinstance(pattern, dict)
    }
    for index, event in enumerate(matrix):
        vehicle_hazard_id = hazard_by_failure.get(event["failure_id"])
        if not isinstance(vehicle_hazard_id, str) or not vehicle_hazard_id:
            raise ValueError(
                f"CS Pack HAZOP 模板缺少 {event['failure_id']} 的 vehicle_hazard_id（事件 {index}）"
            )
        event["vehicle_hazard_id"] = vehicle_hazard_id
        if event.get("severity") == 0:
            # 所有 S=0 事件统一规范化，绝不继承历史模板的 E/C/SG 填法。
            for field in (
                "exposure", "exposure_reason", "controllability",
                "controllability_reason", "safety_goal", "safe_state", "ftti",
            ):
                event[field] = None
            event["expected_asil"] = None

    pack["risk_catalog"]["scenario_sets"]["steering_assist"] = {"scenarios": scenarios}
    pack["risk_catalog"]["event_matrix"] = matrix
    pack["status"] = "cs_risk_matrix_seeded_pending_runtime"
    pack["provenance"]["risk_matrix_source"] = reference.name
    pack["provenance"]["risk_matrix_source_policy"] = (
        "离线从批准基线的语义字段提炼；运行时仅查询 Domain Pack，不读取参考 Excel。"
    )
    if dry_run:
        print(f"[dry-run] 场景={len(scenarios)}，风险矩阵={len(matrix)}，不会写入 {pack_path}")
        return
    pack_path.write_text(json.dumps(pack, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"已写入 {pack_path}: 场景={len(scenarios)}，风险矩阵={len(matrix)}")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("reference", type=Path, help="已批准的 CS 转向 HARA 参考 Excel")
    parser.add_argument("--pack", type=Path, default=DEFAULT_PACK)
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()
    update_pack(args.reference, args.pack, args.dry_run)


if __name__ == "__main__":
    main()
