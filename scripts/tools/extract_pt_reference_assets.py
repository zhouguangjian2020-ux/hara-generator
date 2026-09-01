"""Extract PT reference assets from the five HARA sheets into a format-neutral semantic inventory.

This is an offline migration/audit tool. It reads only the five target sheets and discovers
business columns by header meaning; runtime code never reads the Excel files. Reference
IDs/row positions/merged ranges are retained only as trace evidence, never used as a
matching key.
"""
from __future__ import annotations

import argparse
import json
import re
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
DEFAULT_INPUTS = [
    ROOT / "PT_生产基线_2026-08-21" / "FUSA_PT" / "ALL-HARA_FUSA-100-GPFA-H002______FUSA HARA PT DOMAIN.xlsx",
    ROOT / "PT_生产基线_2026-08-21" / "山子_PT" / "山子高科EE架构项目_PT DOMAIN_HARA_V2.0_12.05.xlsx",
]

FAILURE_MODES = ("丢失", "非预期", "间歇", "过多", "过少", "过早", "反向", "振荡", "部分", "过晚")


def norm(value: Any) -> str:
    if value is None:
        return ""
    return str(value).replace("\u3000", " ").replace("\n", " ").strip()


def header_norm(value: Any) -> str:
    return re.sub(r"[\s\n\r\t/()（）:：\"'“”]", "", norm(value)).lower()


def clean(value: Any) -> Any:
    text = norm(value)
    return text if text else None


def is_marked(value: Any) -> bool:
    return norm(value).lower() in {"√", "✓", "v", "y", "yes", "是", "x", "1"}


def _sheet(wb, predicate, logical_name: str):
    matches = [ws for ws in wb.worksheets if predicate(header_norm(ws.title))]
    if len(matches) != 1:
        raise ValueError(f"{logical_name}: expected one Sheet, got {[ws.title for ws in matches]}")
    return matches[0]


def target_sheets(wb):
    return {
        "s1": _sheet(wb, lambda title: "相关项功能" in title, "S1 related-item feature list"),
        "s2": _sheet(wb, lambda title: "失效模式" in title, "S2 failure modes"),
        "s3": _sheet(wb, lambda title: "hazop" in title and "分析" in title, "S3 HAZOP"),
        "s4": _sheet(wb, lambda title: "hara" in title and "分析" in title, "S4 HARA"),
        "s5": _sheet(wb, lambda title: "整车安全目标" in title, "S5 vehicle safety goals"),
    }


def merged_values(ws) -> dict[tuple[int, int], Any]:
    """Return displayed values; merge propagation only restores visually inherited cells."""
    result: dict[tuple[int, int], Any] = {}
    for merged in ws.merged_cells.ranges:
        anchor = ws.cell(merged.min_row, merged.min_col).value
        for row in range(merged.min_row, merged.max_row + 1):
            for col in range(merged.min_col, merged.max_col + 1):
                result[(row, col)] = anchor
    return result


def val(ws, propagated: dict[tuple[int, int], Any], row: int, col: int) -> Any:
    raw = ws.cell(row, col).value
    return raw if raw is not None else propagated.get((row, col))


def header_columns(ws, propagated: dict[tuple[int, int], Any], rows: Iterable[int] = range(1, 9)) -> dict[int, str]:
    result: dict[int, str] = {}
    for col in range(1, ws.max_column + 1):
        words = []
        for row in rows:
            text = norm(val(ws, propagated, row, col))
            if text and text not in words:
                words.append(text)
        if words:
            result[col] = " ".join(words)
    return result


def choose_column(headers: dict[int, str], logical_name: str, include: tuple[str, ...], *, exclude: tuple[str, ...] = (), required: bool = True) -> int | None:
    candidates = []
    for col, header in headers.items():
        key = header_norm(header)
        if all(header_norm(token) in key for token in include) and not any(header_norm(token) in key for token in exclude):
            candidates.append((len(key), col))
    if candidates:
        return min(candidates)[1]
    if required:
        raise ValueError(f"missing logical column {logical_name}; headers={headers}")
    return None


def rows_with_values(ws, propagated: dict[tuple[int, int], Any], columns: dict[str, int | None], start_row: int = 1):
    for row in range(start_row, ws.max_row + 1):
        yield row, {key: clean(val(ws, propagated, row, col)) if col else None for key, col in columns.items()}


def extract_s1(ws) -> list[dict[str, Any]]:
    p = merged_values(ws); h = header_columns(ws, p, rows=(1,))
    columns = {
        "function_name": choose_column(h, "S1 function", ("整车", "功能"), exclude=("id", "描述", "hara", "章节", "文档")),
        "feature_description": choose_column(h, "S1 feature description", ("功能", "描述")),
        "is_hara": choose_column(h, "S1 is HARA", ("hara",)),
        "chapter": choose_column(h, "S1 chapter", ("章节",), required=False),
        "remark": choose_column(h, "S1 remark", ("备注",), required=False),
    }
    records = []
    for source_row, row in rows_with_values(ws, p, columns):
        if not row["feature_description"] or "功能描述" in row["feature_description"]:
            continue
        records.append({**row, "is_hara": norm(row["is_hara"]) == "是", "source_row": source_row})
    return records


def extract_s2(ws) -> list[dict[str, Any]]:
    p = merged_values(ws); h = header_columns(ws, p, rows=(1,))
    columns = {
        "function_name": choose_column(h, "S2 function", ("功能",), exclude=("id", "理由")),
        "reason": choose_column(h, "S2 reason", ("理由",), required=False),
    }
    modes = {mode: choose_column(h, f"S2 mode {mode}", (mode,), required=False) for mode in FAILURE_MODES}
    records = []
    for source_row, row in rows_with_values(ws, p, columns):
        name = row["function_name"]
        if not name or "功能异常" in name:
            continue
        selected = [mode for mode, col in modes.items() if col and is_marked(val(ws, p, source_row, col))]
        if not selected:
            continue
        records.append({**row, "selected_modes": selected, "source_row": source_row})
    return records


def extract_s3(ws) -> list[dict[str, Any]]:
    p = merged_values(ws); h = header_columns(ws, p, rows=(2,))
    columns = {
        "function_name": choose_column(h, "S3 function", ("整车", "功能"), exclude=("id", "失效模式", "异常")),
        "failure_mode": choose_column(h, "S3 failure mode", ("失效", "模式")),
        "anomaly": choose_column(h, "S3 anomaly", ("异常", "表现")),
        "failure_id": choose_column(h, "S3 failure id", ("失效", "id")),
        "vehicle_hazard": choose_column(h, "S3 vehicle hazard", ("整车", "危害")),
        "related_hara": choose_column(h, "S3 related HARA", ("关联", "hara"), required=False),
        "remark": choose_column(h, "S3 remark", ("备注",), required=False),
    }
    records = []
    for source_row, row in rows_with_values(ws, p, columns):
        if not row["failure_id"] or not row["anomaly"] or "功能失效id" in header_norm(row["failure_id"]):
            continue
        records.append({**row, "source_row": source_row})
    return records


def extract_s4(ws) -> list[dict[str, Any]]:
    p = merged_values(ws); h = header_columns(ws, p, rows=(1, 2))
    columns = {
        "function_name": choose_column(h, "S4 function", ("整车", "功能"), exclude=("id", "失效", "异常")),
        "failure_id": choose_column(h, "S4 failure id", ("失效", "id")),
        "anomaly": choose_column(h, "S4 anomaly", ("异常", "表现")),
        "vehicle_hazard": choose_column(h, "S4 vehicle hazard", ("整车", "危害")),
        "scenario": choose_column(h, "S4 scenario", ("运行", "场景")),
        "event_description": choose_column(h, "S4 event description", ("危害", "事件", "描述")),
        "severity": choose_column(h, "S4 severity", ("严重度",)),
        "severity_reason": choose_column(h, "S4 severity reason", ("s", "理由")),
        "exposure": choose_column(h, "S4 exposure", ("暴露",)),
        "exposure_reason": choose_column(h, "S4 exposure reason", ("e", "理由")),
        "controllability": choose_column(h, "S4 controllability", ("可控性",)),
        "controllability_reason": choose_column(h, "S4 controllability reason", ("可控性", "理由"), required=False),
        "asil": choose_column(h, "S4 ASIL", ("asil",)),
        "safety_goal": choose_column(h, "S4 safety goal", ("安全目标",), exclude=("id",)),
        "safe_state": choose_column(h, "S4 safe state", ("安全状态",)),
        "ftti": choose_column(h, "S4 FTTI", ("ftti",)),
    }
    records = []
    for source_row, row in rows_with_values(ws, p, columns):
        if (not row["failure_id"] or not row["scenario"] or row["severity"] is None
                or "功能失效id" in header_norm(row["failure_id"])):
            continue
        records.append({**row, "source_row": source_row})
    return records


def extract_s5(ws) -> dict[str, list[dict[str, Any]]]:
    p = merged_values(ws)
    # Identify the actual table header semantically.  The title block itself may be merged,
    # so it must not be interpreted as an S5 record.
    header_row = None
    for row in range(1, min(ws.max_row, 12) + 1):
        labels = [header_norm(ws.cell(row, col).value) for col in range(1, ws.max_column + 1)]
        if sum("安全目标id" in label for label in labels) >= 1 and any(label == "asil" for label in labels):
            header_row = row
            break
    if header_row is None:
        raise ValueError("S5 missing semantic table header")

    headers = {col: norm(ws.cell(header_row, col).value) for col in range(1, ws.max_column + 1) if norm(ws.cell(header_row, col).value)}
    id_columns = [col for col, name in headers.items() if "安全目标id" in header_norm(name)]
    if not id_columns:
        raise ValueError(f"S5 missing safety-goal IDs; headers={headers}")

    result = {"goals": [], "merged_goals": []}
    for index, id_col in enumerate(id_columns):
        next_id_col = id_columns[index + 1] if index + 1 < len(id_columns) else ws.max_column + 1
        range_columns = list(range(id_col + 1, next_id_col))
        goal_col = next((col for col in range_columns if "安全目标" in header_norm(headers.get(col, "")) and "id" not in header_norm(headers.get(col, ""))), None)
        asil_col = next((col for col in range_columns if "asil" in header_norm(headers.get(col, ""))), None)
        safe_col = next((col for col in range_columns if "安全状态" in header_norm(headers.get(col, "")) or "safestate" in header_norm(headers.get(col, ""))), None)
        ftti_col = next((col for col in range_columns if "ftti" in header_norm(headers.get(col, ""))), None)
        if not goal_col:
            raise ValueError(f"S5 safety-goal text missing for column {id_col}; headers={headers}")
        target = "goals" if index == 0 else "merged_goals"
        seen: set[tuple[Any, ...]] = set()
        for row in range(header_row + 1, ws.max_row + 1):
            goal_id = clean(val(ws, p, row, id_col))
            goal = clean(val(ws, p, row, goal_col))
            if not goal_id or not goal:
                continue
            item = {
                "safety_goal_id": goal_id, "safety_goal": goal,
                "asil": clean(val(ws, p, row, asil_col)) if asil_col else None,
                "safe_state": clean(val(ws, p, row, safe_col)) if safe_col else None,
                "ftti": clean(val(ws, p, row, ftti_col)) if ftti_col else None,
                "source_row": row,
            }
            signature = tuple(item[key] for key in ("safety_goal_id", "safety_goal", "asil", "safe_state", "ftti"))
            if signature not in seen:
                result[target].append(item)
                seen.add(signature)
    return result

def function_family(name: str | None) -> str:
    text = norm(name)
    if any(token in text for token in ("档位", "换挡")):
        return "pt_gear_state_control"
    if "热管理" in text or "除霜" in text:
        return "pt_thermal_management"
    if "扭矩" in text or "驱动" in text:
        return "pt_torque_control"
    if "充电" in text or "放电" in text or "V2L" in text:
        return "pt_charge_discharge"
    if "高压" in text and any(token in text for token in ("安全", "绝缘", "互锁", "碰撞")):
        return "pt_high_voltage_safety"
    if "高压" in text or "上下电" in text:
        return "pt_high_voltage_state_management"
    if "增程" in text:
        return "pt_range_extender_control"
    return "unclassified"


def project_key(path: Path) -> str:
    return "fusa_pt" if "FUSA" in str(path).upper() else "shanzi_pt"


def extract_book(path: Path) -> dict[str, Any]:
    wb = load_workbook(path, data_only=False, read_only=False)
    try:
        sheets = target_sheets(wb)
        data = {
            "project": project_key(path),
            "source_file": path.name,
            "s1": extract_s1(sheets["s1"]),
            "s2": extract_s2(sheets["s2"]),
            "s3": extract_s3(sheets["s3"]),
            "s4": extract_s4(sheets["s4"]),
            "s5": extract_s5(sheets["s5"]),
        }
    finally:
        wb.close()
    for section in ("s2", "s3"):
        for item in data[section]:
            item["function_family"] = function_family(item.get("function_name"))

    # HARA's "vehicle function" column can be a concrete operating sub-function (for
    # example D-gear control), while S3 contains its parent PT function. Match by the
    # anomaly wording first; failure IDs remain trace-only and are not used here.
    family_by_anomaly: dict[str, str] = {}
    for item in data["s3"]:
        anomaly = norm(item.get("anomaly"))
        family = item.get("function_family")
        if anomaly and family and family != "unclassified":
            family_by_anomaly.setdefault(anomaly, family)
    for item in data["s4"]:
        family = function_family(item.get("function_name"))
        item["function_family"] = family if family != "unclassified" else family_by_anomaly.get(norm(item.get("anomaly")), family)
    return data


def summary(books: list[dict[str, Any]]) -> dict[str, Any]:
    by_family = defaultdict(lambda: {"projects": {}, "s2_mode_sets": {}, "s3": 0, "s4": 0})
    for book in books:
        source = book["project"]
        for s2 in book["s2"]:
            family = s2["function_family"]
            by_family[family]["projects"].setdefault(source, []).append(s2["function_name"])
            by_family[family]["s2_mode_sets"].setdefault(source, []).append(s2["selected_modes"])
        for s3 in book["s3"]:
            by_family[s3["function_family"]]["s3"] += 1
        for s4 in book["s4"]:
            by_family[s4["function_family"]]["s4"] += 1
    normalized = {}
    for family, info in by_family.items():
        per_project_modes = {
            project: sorted({tuple(modes) for modes in lists})
            for project, lists in info["s2_mode_sets"].items()
        }
        mode_signatures = {signature for lists in per_project_modes.values() for signature in lists}
        all_sources_present = len(info["projects"]) == len(books)
        exact_mode_baseline = all_sources_present and len(mode_signatures) == 1
        normalized[family] = {
            "sources": sorted(info["projects"]),
            "source_function_names": {key: sorted(set(values)) for key, values in info["projects"].items()},
            "s2_mode_sets": {key: [list(x) for x in value] for key, value in per_project_modes.items()},
            "s3_anomaly_count": info["s3"],
            "s4_event_count": info["s4"],
            "evidence_status": "exact_known" if exact_mode_baseline else "compatible" if info["projects"] else "unknown",
        }
    return {
        "sources": [book["project"] for book in books],
        "row_counts": {book["project"]: {"s1": len(book["s1"]), "s2": len(book["s2"]), "s3": len(book["s3"]), "s4": len(book["s4"]), "s5_goals": len(book["s5"]["goals"]), "s5_merged": len(book["s5"]["merged_goals"])} for book in books},
        "by_function_family": normalized,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("references", nargs="*", type=Path, default=DEFAULT_INPUTS)
    parser.add_argument("--output", "-o", type=Path, default=ROOT / "PT_reference_assets_extracted_2026-08-21.json")
    args = parser.parse_args()
    books = [extract_book(path.resolve()) for path in args.references]
    output = {"extraction_policy": {"target_sheets_only": True, "semantic_keys_only": True, "source_ids_trace_only": True}, "books": books, "summary": summary(books)}
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(output, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"wrote {args.output}: " + ", ".join(f"{book['project']} S1={len(book['s1'])} S2={len(book['s2'])} S3={len(book['s3'])} S4={len(book['s4'])}" for book in books))


if __name__ == "__main__":
    main()

