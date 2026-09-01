"""
从PT域参考Excel提取案例库
"""
import openpyxl
import json
import sys
from pathlib import Path

def extract_pt_cases(excel_path):
    """提取PT域HARA案例"""
    wb = openpyxl.load_workbook(excel_path, data_only=True)

    # 找到HARA分析表
    hara_sheet = None
    for name in wb.sheetnames:
        if 'HARA' in name.upper() and len(name) < 15:
            hara_sheet = wb[name]
            print(f'Found HARA sheet: {repr(name)}')
            break

    if not hara_sheet:
        print("ERROR: Cannot find HARA sheet")
        return []

    print(f'Total rows: {hara_sheet.max_row}')

    # 列映射（根据观察到的结构）
    COL_HAZARD_ID = 1        # 危害事件ID
    COL_FUNC_NAME = 2        # 功能名称
    COL_FAILURE_ID = 3       # 关联失效ID
    COL_HAZARD_DESC = 4      # 危害识别
    COL_ANOMALY = 4          # 关联异常描述（Row2显示）
    COL_VEHICLE_HAZARD = 5   # 整车危害
    COL_SCENARIO = 6         # 运行场景
    COL_HAZARD_EVENT = 7     # 危害事件描述
    COL_SEVERITY = 8         # 严重度S
    COL_S_REASON = 9         # S理由
    COL_EXPOSURE = 10        # 暴露度E
    COL_E_REASON = 11        # E理由
    COL_CONTROLLABILITY = 12 # 可控性C
    COL_C_REASON = 13        # C理由
    COL_ASIL = 14            # ASIL
    COL_SG_ID = 15           # 安全目标ID
    COL_SAFETY_GOAL = 16     # 安全目标
    COL_SAFE_STATE = 17      # 安全状态
    COL_FTTI = 18            # FTTI
    COL_NOTE = 19            # 注释

    cases = []
    skipped_count = 0
    data_start_row = 3  # 从第3行开始是数据

    for row_idx in range(data_start_row, hara_sheet.max_row + 1):
        # 读取关键字段
        hazard_id = hara_sheet.cell(row_idx, COL_HAZARD_ID).value
        func_name = hara_sheet.cell(row_idx, COL_FUNC_NAME).value
        failure_id = hara_sheet.cell(row_idx, COL_FAILURE_ID).value

        # 跳过空行和分组标题行
        if not hazard_id or not isinstance(hazard_id, str):
            continue
        if not hazard_id.startswith('P_hzrd_'):
            continue

        # 读取完整数据
        anomaly = hara_sheet.cell(row_idx, COL_ANOMALY).value
        vehicle_hazard = hara_sheet.cell(row_idx, COL_VEHICLE_HAZARD).value
        scenario = hara_sheet.cell(row_idx, COL_SCENARIO).value
        hazard_event = hara_sheet.cell(row_idx, COL_HAZARD_EVENT).value
        severity = hara_sheet.cell(row_idx, COL_SEVERITY).value
        s_reason = hara_sheet.cell(row_idx, COL_S_REASON).value
        exposure = hara_sheet.cell(row_idx, COL_EXPOSURE).value
        e_reason = hara_sheet.cell(row_idx, COL_E_REASON).value
        controllability = hara_sheet.cell(row_idx, COL_CONTROLLABILITY).value
        c_reason = hara_sheet.cell(row_idx, COL_C_REASON).value
        asil = hara_sheet.cell(row_idx, COL_ASIL).value
        sg_id = hara_sheet.cell(row_idx, COL_SG_ID).value
        safety_goal = hara_sheet.cell(row_idx, COL_SAFETY_GOAL).value
        safe_state = hara_sheet.cell(row_idx, COL_SAFE_STATE).value
        ftti = hara_sheet.cell(row_idx, COL_FTTI).value
        note = hara_sheet.cell(row_idx, COL_NOTE).value

        # 数据清洗
        def clean_str(val):
            if val is None:
                return None
            s = str(val).strip()
            return s if s else None

        def clean_int(val):
            if val is None or val == '':
                return None
            try:
                return int(val)
            except:
                return None

        # 提取分析单元（从功能名称）
        analysis_unit = clean_str(func_name)

        # 提取失效模式（从failure_id）
        failure_mode = None
        if failure_id:
            # P_MF_0001_01 -> 从HAZOP表推断失效模式
            # 这里先保留failure_id，后续关联
            failure_mode = clean_str(failure_id)

        # 构建案例
        case = {
            'case_id': hazard_id,
            'domain': 'PT',
            'analysis_unit': analysis_unit,
            'failure_id': failure_mode,
            'anomaly': clean_str(anomaly),
            'vehicle_hazard': clean_str(vehicle_hazard),
            'scenario': clean_str(scenario),
            'hazard_event_desc': clean_str(hazard_event),
            'severity': clean_int(severity),
            'severity_reason': clean_str(s_reason),
            'exposure': clean_int(exposure),
            'exposure_reason': clean_str(e_reason),
            'controllability': clean_int(controllability),
            'controllability_reason': clean_str(c_reason),
            'asil': clean_str(asil),
            'safety_goal_id': clean_str(sg_id),
            'safety_goal': clean_str(safety_goal),
            'safe_state': clean_str(safe_state),
            'ftti': clean_str(ftti),
            'note': clean_str(note)
        }

        # 验证必要字段
        if not case['anomaly'] or not case['scenario']:
            skipped_count += 1
            continue

        # S=0的特殊处理
        if case['severity'] == 0:
            # S=0时，E/C/ASIL/SG应该为空
            if case['exposure'] is not None or case['controllability'] is not None:
                print(f"WARNING Row {row_idx}: S=0 but E/C not empty: {hazard_id}")
            case['exposure'] = None
            case['exposure_reason'] = None
            case['controllability'] = None
            case['controllability_reason'] = None
            case['asil'] = None
            case['safety_goal'] = None
            case['safety_goal_id'] = None
            case['safe_state'] = None
            case['ftti'] = None

        cases.append(case)

    print(f'\nExtracted {len(cases)} valid cases')
    print(f'Skipped {skipped_count} incomplete rows')

    return cases


def analyze_cases(cases):
    """分析案例库统计信息"""
    print('\n=== 案例库统计 ===')
    print(f'总案例数: {len(cases)}')

    # 按分析单元统计
    units = {}
    for case in cases:
        unit = case['analysis_unit']
        if unit not in units:
            units[unit] = []
        units[unit].append(case)

    print(f'\n分析单元数: {len(units)}')
    for unit, unit_cases in sorted(units.items(), key=lambda x: len(x[1]), reverse=True):
        print(f'  {unit}: {len(unit_cases)} 案例')

    # 按场景统计
    scenarios = {}
    for case in cases:
        scenario = case['scenario']
        if scenario:
            scenarios[scenario] = scenarios.get(scenario, 0) + 1

    print(f'\n场景数: {len(scenarios)}')
    top_scenarios = sorted(scenarios.items(), key=lambda x: x[1], reverse=True)[:10]
    for scenario, count in top_scenarios:
        print(f'  {scenario}: {count} 次')

    # ASIL分布
    asil_dist = {}
    for case in cases:
        asil = case['asil'] or 'NULL'
        asil_dist[asil] = asil_dist.get(asil, 0) + 1

    print(f'\nASIL分布:')
    for asil, count in sorted(asil_dist.items()):
        print(f'  {asil}: {count}')

    # S=0统计
    s0_count = sum(1 for c in cases if c['severity'] == 0)
    print(f'\nS=0 案例: {s0_count}')

    # 检查数据完整性
    print(f'\n=== 数据完整性检查 ===')
    missing_fields = {
        'anomaly': 0,
        'vehicle_hazard': 0,
        'scenario': 0,
        'severity': 0,
        'exposure': 0,
        'controllability': 0
    }

    for case in cases:
        if case['severity'] == 0:
            # S=0不检查E/C
            continue
        for field in missing_fields:
            if case[field] is None:
                missing_fields[field] += 1

    for field, count in missing_fields.items():
        if count > 0:
            print(f'  缺失{field}: {count} 案例')

    # 检查S/E/C范围
    invalid_sec = []
    for case in cases:
        if case['severity'] == 0:
            continue
        s, e, c = case['severity'], case['exposure'], case['controllability']
        if s is not None and (s < 0 or s > 3):
            invalid_sec.append(f"{case['case_id']}: S={s}")
        if e is not None and (e < 0 or e > 4):
            invalid_sec.append(f"{case['case_id']}: E={e}")
        if c is not None and (c < 0 or c > 3):
            invalid_sec.append(f"{case['case_id']}: C={c}")

    if invalid_sec:
        print(f'\n无效S/E/C值:')
        for msg in invalid_sec[:5]:
            print(f'  {msg}')
        if len(invalid_sec) > 5:
            print(f'  ... 还有 {len(invalid_sec) - 5} 个')
    else:
        print(f'\n[OK] 所有S/E/C值有效')


def main():
    excel_path = Path('关键数据/PT域/ALL-HARA_FUSA-100-GPFA-H002______FUSA HARA PT DOMAIN.xlsx')

    if not excel_path.exists():
        print(f'ERROR: File not found: {excel_path}')
        sys.exit(1)

    print(f'Extracting cases from: {excel_path}')
    cases = extract_pt_cases(excel_path)

    if not cases:
        print('ERROR: No cases extracted')
        sys.exit(1)

    # 统计分析
    analyze_cases(cases)

    # 保存为JSON
    output_path = Path('关键数据/PT域/PT_extracted_cases.json')
    with open(output_path, 'w', encoding='utf-8') as f:
        json.dump(cases, f, ensure_ascii=False, indent=2)

    print(f'\n[OK] Cases saved to: {output_path}')

    # 生成简化预览
    preview_path = Path('关键数据/PT域/PT_cases_preview.txt')
    with open(preview_path, 'w', encoding='utf-8') as f:
        f.write('PT域案例库预览\n')
        f.write('=' * 80 + '\n\n')
        for i, case in enumerate(cases[:20], 1):
            f.write(f'{i}. {case["case_id"]}\n')
            f.write(f'   分析单元: {case["analysis_unit"]}\n')
            f.write(f'   异常: {case["anomaly"]}\n')
            f.write(f'   场景: {case["scenario"]}\n')
            f.write(f'   S/E/C: {case["severity"]}/{case["exposure"]}/{case["controllability"]}\n')
            f.write(f'   ASIL: {case["asil"]}\n')
            f.write(f'   安全目标: {case["safety_goal"]}\n')
            f.write('\n')
        f.write(f'... 共 {len(cases)} 个案例\n')

    print(f'[OK] Preview saved to: {preview_path}')


if __name__ == '__main__':
    main()
