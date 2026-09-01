# extract_s2_from_excel.py
# 从 11 个参考 Excel 中提取失效模式 Sheet 数据，输出 JSON 查表
# 用法: python scripts/tools/extract_s2_from_excel.py

import json
import sys
from pathlib import Path

import openpyxl

FAILURE_MODES = [
    "丢失", "非预期", "间歇", "过多", "过少",
    "过早", "反向", "振荡", "部分", "过晚", "卡滞",
]

# 11 个数据组
DATA_GROUPS = [
    {"group": 1, "domain": "CS", "path": r"R:\Users\zhouguangjian\Developer\projects\数据组\数据组1\山子高科EE架构项目_STEERING_HARA_V2.0_12.10.xlsx"},
    {"group": 2, "domain": "CB", "path": r"R:\Users\zhouguangjian\Developer\projects\数据组\数据组2\山子高科EE架构项目_BRAKE EPB_HARA_V2.0_12.30.xlsx"},
    {"group": 3, "domain": "CB", "path": r"R:\Users\zhouguangjian\Developer\projects\数据组\数据组3\山子高科EE架构项目_BRAKE_HARA_V2.0_12.02.xlsx"},
    {"group": 4, "domain": "P",  "path": r"R:\Users\zhouguangjian\Developer\projects\数据组\数据组4\山子高科EE架构项目_PT DOMAIN_HARA_V2.0_12.05.xlsx"},
    {"group": 5, "domain": "P",  "path": r"R:\Users\zhouguangjian\Developer\projects\数据组\数据组5\ALL-HARA_FUSA-100-GPFA-H002______FUSA HARA PT DOMAIN.xlsx"},
    {"group": 6, "domain": "CB", "path": r"R:\Users\zhouguangjian\Developer\projects\数据组\数据组6\ALL-HARA_FUSA-200-GPFA-H002______FUSA HARA CH BRAKE.xlsx"},
    {"group": 7, "domain": "CB", "path": r"R:\Users\zhouguangjian\Developer\projects\数据组\数据组7\ALL-HARA_FUSA-200-GPFA-H102______FUSA HARA CH BRAKE EPB.xlsx"},
    {"group": 8, "domain": "CS", "path": r"R:\Users\zhouguangjian\Developer\projects\数据组\数据组8\ALL-HARA_FUSA-210-GPFA-H002______FUSA HARA CH STEERING.xlsx"},
    {"group": 9, "domain": "A",  "path": r"R:\Users\zhouguangjian\Developer\projects\数据组\数据组9\ALL-HARA_FUSA-300-GPFA-H002______FUSA HARA AD DOMAIN.xlsx"},
    {"group": 10,"domain": "Info","path": r"R:\Users\zhouguangjian\Developer\projects\数据组\数据组10\ALL-HARA_FUSA-400-GPFA-H002______FUSA HARA ET DOMAIN.xlsx"},
    {"group": 11,"domain": "B",  "path": r"R:\Users\zhouguangjian\Developer\projects\数据组\数据组11\ALL-HARA_FUSA-500-GPFA-H002______FUSA HARA BD DOMAIN.xlsx"},
]


def find_failure_mode_sheet(wb):
    """在 workbook 中找失效模式 Sheet"""
    for name in wb.sheetnames:
        if "失效模式" in name:
            return wb[name]
    for name in wb.sheetnames:
        if "失效" in name:
            return wb[name]
    return None


def extract_from_sheet(ws):
    """从失效模式 Sheet 提取数据行"""
    mode_cols = {}
    id_col = None
    name_col = None
    reason_col = None

    for col in range(1, ws.max_column + 1):
        header = ws.cell(row=1, column=col).value
        if not header:
            continue
        h = str(header).strip()
        if h in FAILURE_MODES:
            mode_cols[h] = col
        if "整车功能ID" in h or "整车功能 ID" in h or h == "整车功能ID":
            id_col = col
        if "功能异常" in h or "整车功能" == h or "功能异常\n整车功能" == h:
            name_col = col
        if "选择理由" in h or "理由" in h:
            reason_col = col

    if name_col is None:
        name_col = 2

    rows = []
    for r in range(2, ws.max_row + 1):
        func_name_raw = ws.cell(row=r, column=name_col).value
        if not func_name_raw:
            continue
        func_name = str(func_name_raw).strip()
        if not func_name or func_name.startswith("注") or func_name.startswith("说明"):
            continue

        selected = []
        for mode, col in mode_cols.items():
            cell_val = ws.cell(row=r, column=col).value
            if cell_val and str(cell_val).strip() in ["√", "√", "V", "v", "✓", "Y", "y", "是"]:
                selected.append(mode)

        reason = ""
        if reason_col:
            reason_val = ws.cell(row=r, column=reason_col).value
            if reason_val:
                reason = str(reason_val).strip()

        func_id = ""
        if id_col:
            fid_val = ws.cell(row=r, column=id_col).value
            if fid_val:
                func_id = str(fid_val).strip()

        rows.append({
            "func_id": func_id,
            "func_name": func_name,
            "selected_modes": selected,
            "reason": reason,
            "available_modes": list(mode_cols.keys()),
            "missing_modes": [m for m in FAILURE_MODES if m not in mode_cols],
        })

    return rows


def main():
    all_entries = []

    for dg in DATA_GROUPS:
        path = dg["path"]
        if not Path(path).exists():
            print(f"[WARN] 文件不存在: {path}")
            continue

        print(f"\n=== 数据组{dg['group']} ({dg['domain']}) ===")
        print(f"  文件: {Path(path).name}")

        try:
            wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        except Exception as e:
            print(f"  [ERROR] 无法打开: {e}")
            continue

        ws = find_failure_mode_sheet(wb)
        if ws is None:
            print(f"  [WARN] 未找到失效模式 Sheet")
            print(f"  Sheet 列表: {wb.sheetnames}")
            wb.close()
            continue

        print(f"  Sheet: {ws.title}")
        print(f"  尺寸: {ws.max_row} 行 x {ws.max_column} 列")

        rows = extract_from_sheet(ws)
        print(f"  提取到 {len(rows)} 行数据")
        for r in rows:
            print(f"    {r['func_id']:20s} | {r['func_name']:30s} | 模式: {','.join(r['selected_modes']):40s} | 缺列: {','.join(r['missing_modes']) if r['missing_modes'] else '无'}")
            all_entries.append({
                "group": dg["group"],
                "domain": dg["domain"],
                "func_id": r["func_id"],
                "func_name": r["func_name"],
                "selected_modes": r["selected_modes"],
                "reason": r["reason"],
                "missing_modes": r["missing_modes"],
                "source": "reference_excel",
            })

        wb.close()

    output_path = Path(__file__).resolve().parent.parent.parent / "references" / "s2-lookup-table.json"
    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(all_entries, f, ensure_ascii=False, indent=2)
    print(f"\n=== 完成 ===")
    print(f"共 {len(all_entries)} 条记录")
    print(f"输出: {output_path}")


if __name__ == "__main__":
    main()
