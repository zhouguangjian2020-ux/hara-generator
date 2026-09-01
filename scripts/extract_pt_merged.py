#!/usr/bin/env python3
"""
合并两个PT域Excel文件的数据到一个完整的Domain Pack
Excel1: ALL-HARA_FUSA-100-GPFA-H002______FUSA HARA PT DOMAIN.xlsx (451条HARA)
Excel2: FUSA HARA PT DOMAIN_V1.1_山子高科 - 副本.xlsx (34条HARA)
"""

import json
import openpyxl
from pathlib import Path
from collections import defaultdict
import sys

def get_cell_value(ws, row, col):
    """获取单元格值，处理合并单元格"""
    cell = ws.cell(row, col)
    for merged_range in ws.merged_cells.ranges:
        if cell.coordinate in merged_range:
            top_left_cell = ws.cell(merged_range.min_row, merged_range.min_col)
            return top_left_cell.value
    return cell.value

def normalize_text(text):
    """标准化文本"""
    if not text or not isinstance(text, str):
        return ""
    return " ".join(text.strip().split())

def extract_failure_modes(wb):
    """提取失效模式表"""
    if '失效模式' not in wb.sheetnames:
        return []

    ws = wb['失效模式']
    print(f"  失效模式表: {ws.max_row}行")

    # 第1行是表头
    headers = {}
    for col in range(1, min(ws.max_column + 1, 20)):
        val = get_cell_value(ws, 1, col)
        if val:
            headers[col] = normalize_text(str(val))

    # 识别列
    col_func_id = None
    col_func_name = None
    col_reason = None
    failure_mode_cols = {}

    for col, header in headers.items():
        if "功能" in header and "ID" in header:
            col_func_id = col
        elif "功能" in header:
            col_func_name = col
        elif "理由" in header:
            col_reason = col
        elif header in ["丢失", "非预期", "间歇", "过早", "过晚", "过多", "过少", "反向", "振荡", "部分", "卡滞"]:
            failure_mode_cols[col] = header

    # 提取数据
    records = []
    for row in range(2, ws.max_row + 1):
        func_id = get_cell_value(ws, row, col_func_id) if col_func_id else None

        # 跳过注释行
        if func_id and "注：" in str(func_id):
            break

        if not func_id:
            continue

        func_name = get_cell_value(ws, row, col_func_name) if col_func_name else None
        if not func_name:
            continue

        func_id = str(func_id).strip()
        func_name = normalize_text(func_name)
        reason = normalize_text(get_cell_value(ws, row, col_reason)) if col_reason else ""

        # 检查选中的失效模式
        selected_modes = []
        for col, mode_name in failure_mode_cols.items():
            val = get_cell_value(ws, row, col)
            if val and str(val).strip() in ["√", "✓", "是", "Y", "y", "1", "有"]:
                selected_modes.append(mode_name)

        if selected_modes:
            records.append({
                "function_id": func_id,
                "function_name": func_name,
                "failure_modes": selected_modes,
                "reason": reason
            })

    print(f"    提取到 {len(records)} 个功能")
    return records

def extract_hazop(wb):
    """提取HAZOP表"""
    if 'HAZOP 分析' not in wb.sheetnames:
        return []

    ws = wb['HAZOP 分析']
    print(f"  HAZOP表: {ws.max_row}行")

    # 查找表头行
    header_row = 2
    for row in range(1, 5):
        val = get_cell_value(ws, row, 1)
        if val and "功能" in str(val) and "ID" in str(val):
            header_row = row
            break

    headers = {}
    for col in range(1, min(ws.max_column + 1, 20)):
        val = get_cell_value(ws, header_row, col)
        if val:
            headers[col] = normalize_text(str(val))

    # 识别列
    col_map = {}
    for col, header in headers.items():
        if "功能" in header and "ID" in header and "失效" not in header:
            col_map["function_id"] = col
        elif "整车功能" in header and "ID" not in header:
            col_map["function"] = col
        elif "失效模式" in header:
            col_map["failure_mode"] = col
        elif "异常表现" in header:
            col_map["anomaly"] = col
        elif "失效" in header and "ID" in header:
            col_map["failure_id"] = col
        elif "整车危害" in header:
            col_map["vehicle_hazard"] = col

    # 提取数据
    records = []
    current_func_id = None
    current_func = None

    for row in range(header_row + 1, ws.max_row + 1):
        func_id = get_cell_value(ws, row, col_map.get("function_id", 1))
        func = get_cell_value(ws, row, col_map.get("function", 2))

        if func_id:
            current_func_id = str(func_id).strip()
        if func:
            current_func = normalize_text(func)

        failure_mode = get_cell_value(ws, row, col_map.get("failure_mode", 3))
        if not failure_mode:
            continue

        record = {
            "function_id": current_func_id if current_func_id else "",
            "function": current_func if current_func else "",
            "failure_mode": normalize_text(failure_mode),
            "anomaly": normalize_text(get_cell_value(ws, row, col_map.get("anomaly", 4))),
            "failure_id": str(get_cell_value(ws, row, col_map.get("failure_id", 5)) or "").strip(),
            "vehicle_hazard": normalize_text(get_cell_value(ws, row, col_map.get("vehicle_hazard", 6))),
        }
        records.append(record)

    print(f"    提取到 {len(records)} 条记录")
    return records

def extract_hara(wb):
    """提取HARA表"""
    if 'HARA分析  ' not in wb.sheetnames:
        return []

    ws = wb['HARA分析  ']
    print(f"  HARA表: {ws.max_row}行")

    # 查找表头行（有些表头在第2行，有些合并了第1-2行）
    header_row = 2
    for row in range(1, 5):
        for col in range(1, 10):
            val = get_cell_value(ws, row, col)
            if val and "严重度" in str(val):
                header_row = row
                break

    headers = {}
    for col in range(1, min(ws.max_column + 1, 25)):
        val = get_cell_value(ws, header_row, col)
        if val:
            headers[col] = normalize_text(str(val))

    # 识别列
    col_map = {}
    for col, header in headers.items():
        h_lower = header.lower()
        if "危害事件" in header and "id" in h_lower:
            col_map["hazard_id"] = col
        elif "整车功能" in header and "id" not in h_lower:
            col_map["function"] = col
        elif "功能失效" in header and "id" in h_lower:
            col_map["failure_id"] = col
        elif "异常表现" in header:
            col_map["anomaly"] = col
        elif "整车危害" in header and "id" not in h_lower:
            col_map["vehicle_hazard"] = col
        elif "运行场景" in header or "驾驶" in header:
            col_map["scenario"] = col
        elif "危害事件描述" in header:
            col_map["description"] = col
        elif "严重度" in header:
            col_map["severity"] = col
        elif "s" in h_lower and "理由" in header:
            col_map["severity_reason"] = col
        elif "暴露" in header:
            col_map["exposure"] = col
        elif "e" in h_lower and "理由" in header:
            col_map["exposure_reason"] = col
        elif "可控" in header:
            col_map["controllability"] = col
        elif "c" in h_lower and "理由" in header:
            col_map["controllability_reason"] = col
        elif header.strip().upper() == "ASIL":
            col_map["asil"] = col
        elif "安全目标" in header and "id" in h_lower:
            col_map["safety_goal_id"] = col
        elif "安全目标" in header and "id" not in h_lower:
            col_map["safety_goal"] = col
        elif "安全状态" in header:
            col_map["safe_state"] = col
        elif "ftti" in h_lower:
            col_map["ftti"] = col

    # 提取数据
    records = []
    current_hazard_id = None
    current_function = None
    current_failure_id = None

    for row in range(header_row + 1, ws.max_row + 1):
        anomaly = get_cell_value(ws, row, col_map.get("anomaly", 4))
        vehicle_hazard = get_cell_value(ws, row, col_map.get("vehicle_hazard", 5))

        if not anomaly and not vehicle_hazard:
            continue

        hazard_id = get_cell_value(ws, row, col_map.get("hazard_id", 1))
        function = get_cell_value(ws, row, col_map.get("function", 2))
        failure_id = get_cell_value(ws, row, col_map.get("failure_id", 3))

        if hazard_id:
            current_hazard_id = str(hazard_id).strip()
        if function:
            current_function = normalize_text(function)
        if failure_id:
            current_failure_id = str(failure_id).strip()

        def parse_sec(val):
            if val is None:
                return None
            val_str = str(val).strip()
            if val_str in ("", "-", "N/A", "n/a", "NA"):
                return None
            try:
                return int(float(val_str))
            except:
                return None

        severity = parse_sec(get_cell_value(ws, row, col_map.get("severity", 8)))
        exposure = parse_sec(get_cell_value(ws, row, col_map.get("exposure", 10)))
        controllability = parse_sec(get_cell_value(ws, row, col_map.get("controllability", 12)))

        asil_val = get_cell_value(ws, row, col_map.get("asil", 14))
        asil = str(asil_val).strip().upper() if asil_val and str(asil_val).strip() not in ("", "-", "N/A") else None

        record = {
            "hazard_id": current_hazard_id if current_hazard_id else "",
            "function": current_function if current_function else "",
            "failure_id": current_failure_id if current_failure_id else "",
            "anomaly": normalize_text(anomaly) if anomaly else "",
            "vehicle_hazard": normalize_text(vehicle_hazard) if vehicle_hazard else "",
            "scenario": normalize_text(get_cell_value(ws, row, col_map.get("scenario", 6))),
            "description": normalize_text(get_cell_value(ws, row, col_map.get("description", 7))),
            "severity": severity,
            "severity_reason": normalize_text(get_cell_value(ws, row, col_map.get("severity_reason", 9))),
            "exposure": exposure,
            "exposure_reason": normalize_text(get_cell_value(ws, row, col_map.get("exposure_reason", 11))),
            "controllability": controllability,
            "controllability_reason": normalize_text(get_cell_value(ws, row, col_map.get("controllability_reason", 13))),
            "asil": asil,
            "safety_goal_id": str(get_cell_value(ws, row, col_map.get("safety_goal_id", 15)) or "").strip(),
            "safety_goal": normalize_text(get_cell_value(ws, row, col_map.get("safety_goal", 16))),
            "safe_state": normalize_text(get_cell_value(ws, row, col_map.get("safe_state", 17))),
            "ftti": normalize_text(get_cell_value(ws, row, col_map.get("ftti", 18))),
        }
        records.append(record)

    print(f"    提取到 {len(records)} 条记录")
    return records

def extract_hazard_family(vehicle_hazard):
    """从整车危害中提取危害族"""
    if not vehicle_hazard:
        return "unknown"

    vh = vehicle_hazard.lower()

    if any(w in vh for w in ["加速", "减速", "制动", "纵向", "速度", "扭矩"]):
        if "非预期" in vh:
            return "unexpected_longitudinal_motion"
        elif "丢失" in vh or "失效" in vh:
            return "loss_of_longitudinal_control"
        return "longitudinal_motion"

    if any(w in vh for w in ["转向", "横向", "方向"]):
        if "非预期" in vh:
            return "unexpected_lateral_motion"
        return "lateral_motion"

    if any(w in vh for w in ["能量", "电", "充电"]):
        return "energy_management"

    if any(w in vh for w in ["显示", "信息"]):
        return "information_display"

    return "other"

def merge_data(all_failure_modes, all_hazop, all_hara):
    """合并数据并构建Domain Pack"""
    print(f"\n{'='*60}")
    print("合并数据并构建Domain Pack")
    print(f"{'='*60}")

    # 1. 合并失效模式（去重）
    fm_dict = {}
    for fm in all_failure_modes:
        func_id = fm["function_id"]
        if func_id not in fm_dict:
            fm_dict[func_id] = fm["failure_modes"]
        else:
            # 合并失效模式列表
            existing = set(fm_dict[func_id])
            existing.update(fm["failure_modes"])
            fm_dict[func_id] = list(existing)

    print(f"失效模式规则: {len(fm_dict)} 个功能")

    # 2. HAZOP模式（去重）
    hazop_patterns = []
    seen_patterns = set()
    for haz in all_hazop:
        if haz["failure_id"] and haz["failure_mode"]:
            key = (haz["failure_id"], haz["failure_mode"])
            if key not in seen_patterns:
                seen_patterns.add(key)
                pattern = {
                    "failure_mode": haz["failure_mode"],
                    "failure_id": haz["failure_id"],
                    "anomaly": haz["anomaly"],
                    "hazard_family": extract_hazard_family(haz["vehicle_hazard"]),
                    "vehicle_hazard": haz["vehicle_hazard"]
                }
                hazop_patterns.append(pattern)

    print(f"HAZOP模式: {len(hazop_patterns)} 条")

    # 3. 场景（去重）
    scenarios = {}
    counter = 1
    for rec in all_hara:
        if rec["scenario"] and rec["scenario"] not in scenarios:
            scenarios[rec["scenario"]] = f"PT_SC_{counter:03d}"
            counter += 1

    print(f"场景: {len(scenarios)} 个")

    # 4. 事件矩阵
    event_matrix = []
    for rec in all_hara:
        scenario_id = scenarios.get(rec["scenario"], "PT_SC_999")
        event = {
            "analysis_unit_id": "PT_MAIN",
            "failure_id": rec["failure_id"],
            "scenario_id": scenario_id,
            "description": rec["description"],
            "severity": rec["severity"],
            "severity_reason": rec["severity_reason"],
            "exposure": rec["exposure"],
            "exposure_reason": rec["exposure_reason"],
            "controllability": rec["controllability"],
            "controllability_reason": rec["controllability_reason"],
            "expected_asil": rec["asil"],
            "safety_goal": rec["safety_goal"] if rec["safety_goal"] else None,
            "safe_state": rec["safe_state"] if rec["safe_state"] else None,
            "ftti": rec["ftti"] if rec["ftti"] else None,
            "vehicle_hazard_id": rec["hazard_id"],
        }
        event_matrix.append(event)

    print(f"事件矩阵: {len(event_matrix)} 条")

    # 5. 安全目标目录
    hazard_families = {}
    hazard_catalog = {}

    for rec in all_hara:
        if rec["hazard_id"] and rec["vehicle_hazard"]:
            hazard_family = extract_hazard_family(rec["vehicle_hazard"])

            hazard_catalog[rec["hazard_id"]] = {
                "hazard_family": hazard_family,
                "vehicle_hazard": rec["vehicle_hazard"]
            }

            if hazard_family not in hazard_families and rec["safety_goal"]:
                hazard_families[hazard_family] = {
                    "canonical_safety_goal": rec["safety_goal"],
                    "safe_state": rec["safe_state"] if rec["safe_state"] else None,
                    "ftti": rec["ftti"] if rec["ftti"] else None,
                    "canonical_hazard": rec["vehicle_hazard"]
                }

    print(f"危害族: {len(hazard_families)} 个")
    print(f"危害目录: {len(hazard_catalog)} 条")

    domain_pack = {
        "schema_version": "1.0",
        "domain": "PT",
        "domain_name": "动力域",
        "status": "pt_complete_merged",
        "provenance": {
            "migration_stage": "P0-1",
            "extraction_date": "2026-08-26",
            "source_files": [
                "ALL-HARA_FUSA-100-GPFA-H002______FUSA HARA PT DOMAIN.xlsx",
                "FUSA HARA PT DOMAIN_V1.1_山子高科 - 副本.xlsx"
            ],
            "merge_policy": "两个Excel文件数据合并，失效模式、HAZOP、场景去重"
        },
        "function_catalog": {
            "function_matchers": [],
            "semantic_roles": {},
            "scope_rules": []
        },
        "analysis_catalog": {
            "analysis_units": [],
            "failure_mode_rules": fm_dict,
            "hazop_patterns": hazop_patterns
        },
        "risk_catalog": {
            "scenario_sets": {
                "pt_main": {
                    "scenarios": [
                        {"scenario_id": sid, "text": stext}
                        for stext, sid in sorted(scenarios.items(), key=lambda x: x[1])
                    ]
                }
            },
            "scenario_lookup": {},
            "event_matrix": event_matrix
        },
        "safety_goal_catalog": {
            "hazard_families": hazard_families,
            "rules": [{
                "deduplication_key": ["domain", "hazard_family", "canonical_safety_goal", "safe_state", "ftti"]
            }],
            "hazard_catalog": hazard_catalog
        }
    }

    return domain_pack

def main():
    sys.stdout.reconfigure(encoding='utf-8')

    excel1_path = r"E:\Git projects\hara-generator-skill\数据组\数据组5\ALL-HARA_FUSA-100-GPFA-H002______FUSA HARA PT DOMAIN.xlsx"
    excel2_path = r"C:\Users\zhouguangjian\Downloads\FUSA HARA PT DOMAIN_V1.1_山子高科 - 副本.xlsx"

    all_failure_modes = []
    all_hazop = []
    all_hara = []

    # Excel 1
    print(f"{'='*60}")
    print(f"提取Excel文件1")
    print(f"{'='*60}")
    wb1 = openpyxl.load_workbook(excel1_path, data_only=True)
    all_failure_modes.extend(extract_failure_modes(wb1))
    all_hazop.extend(extract_hazop(wb1))
    all_hara.extend(extract_hara(wb1))

    # Excel 2
    print(f"\n{'='*60}")
    print(f"提取Excel文件2")
    print(f"{'='*60}")
    wb2 = openpyxl.load_workbook(excel2_path, data_only=True)
    all_failure_modes.extend(extract_failure_modes(wb2))
    all_hazop.extend(extract_hazop(wb2))
    all_hara.extend(extract_hara(wb2))

    # 合并
    domain_pack = merge_data(all_failure_modes, all_hazop, all_hara)

    # 保存
    output_path = Path("references/domain_packs/PT_complete.json")
    output_path.parent.mkdir(parents=True, exist_ok=True)

    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(domain_pack, f, ensure_ascii=False, indent=2)

    print(f"\n完整Domain Pack已保存到: {output_path}")

    # 统计ASIL分布
    asil_dist = defaultdict(int)
    for event in domain_pack['risk_catalog']['event_matrix']:
        asil = event.get("expected_asil")
        asil_dist[asil if asil else "null"] += 1

    print(f"\nASIL分布:")
    for asil, count in sorted(asil_dist.items(), key=lambda x: str(x[0])):
        print(f"  {asil}: {count}")

if __name__ == "__main__":
    main()
