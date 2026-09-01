#!/usr/bin/env python3
"""
从PT域参考Excel提取数据并生成Domain Pack格式的预填表
特别处理合并单元格的情况
"""

import json
import openpyxl
from pathlib import Path
from collections import defaultdict
import sys

def get_cell_value(ws, row, col):
    """
    获取单元格值，处理合并单元格的情况
    如果是合并单元格的一部分，返回合并区域左上角的值
    """
    cell = ws.cell(row, col)

    # 检查是否是合并单元格
    for merged_range in ws.merged_cells.ranges:
        if cell.coordinate in merged_range:
            # 返回合并区域左上角单元格的值
            top_left_cell = ws.cell(merged_range.min_row, merged_range.min_col)
            return top_left_cell.value

    return cell.value

def normalize_text(text):
    """标准化文本：去除多余空格和换行"""
    if not text or not isinstance(text, str):
        return ""
    return " ".join(text.strip().split())

def find_header_row(ws, max_search_rows=10):
    """
    查找表头行：寻找包含关键字段的行
    """
    key_headers = ["功能异常", "整车危害", "运行场景", "严重度", "暴露", "可控", "ASIL"]

    for row in range(1, max_search_rows + 1):
        found_count = 0
        for col in range(1, ws.max_column + 1):
            cell_value = get_cell_value(ws, row, col)
            if cell_value:
                cell_text = str(cell_value)
                for key in key_headers:
                    if key in cell_text:
                        found_count += 1
                        break

        if found_count >= 4:  # 至少找到4个关键字段
            return row

    return 2  # 默认第2行

def extract_hara_sheet(ws):
    """
    提取HARA分析表数据，特别处理合并单元格
    """
    print(f"\n开始提取HARA表，总行数: {ws.max_row}, 总列数: {ws.max_column}")

    # 打印合并单元格信息
    print(f"合并单元格数量: {len(ws.merged_cells.ranges)}")
    if len(ws.merged_cells.ranges) > 0:
        print("前10个合并区域:")
        for i, merged_range in enumerate(list(ws.merged_cells.ranges)[:10]):
            print(f"  {i+1}. {merged_range}")

    # 查找表头行
    header_row = find_header_row(ws)
    print(f"\n表头行位于第 {header_row} 行")

    # 读取表头
    headers = {}
    for col in range(1, ws.max_column + 1):
        cell_value = get_cell_value(ws, header_row, col)
        if cell_value:
            headers[col] = normalize_text(str(cell_value))

    print(f"\n表头内容:")
    for col, header in sorted(headers.items()):
        print(f"  列{col}: {header}")

    # 定位关键列（模糊匹配）
    col_map = {}
    for col, header in headers.items():
        header_lower = header.lower()

        if "危害事件" in header and "id" in header_lower:
            col_map["hazard_id"] = col
        elif "整车功能" in header and "id" not in header_lower:
            col_map["function"] = col
        elif "功能失效" in header and "id" in header_lower:
            col_map["failure_id"] = col
        elif "功能异常" in header or "异常表现" in header:
            col_map["anomaly"] = col
        elif "整车危害" in header and "id" not in header_lower:
            col_map["vehicle_hazard"] = col
        elif "运行场景" in header or "驾驶" in header:
            col_map["scenario"] = col
        elif "危害事件描述" in header or ("描述" in header and "危害" in header):
            col_map["description"] = col
        elif "严重度" in header or (header.startswith("s") and len(header) <= 3):
            col_map["severity"] = col
        elif "s" in header_lower and "理由" in header:
            col_map["severity_reason"] = col
        elif "暴露" in header or "e" in header_lower and "概率" in header:
            col_map["exposure"] = col
        elif "e" in header_lower and "理由" in header:
            col_map["exposure_reason"] = col
        elif "可控" in header or (header.startswith("c") and len(header) <= 5):
            col_map["controllability"] = col
        elif "c" in header_lower and "理由" in header:
            col_map["controllability_reason"] = col
        elif header.strip().upper() == "ASIL":
            col_map["asil"] = col
        elif "安全目标" in header and "id" in header_lower:
            col_map["safety_goal_id"] = col
        elif "安全目标" in header and "id" not in header_lower:
            col_map["safety_goal"] = col
        elif "安全状态" in header:
            col_map["safe_state"] = col
        elif "ftti" in header_lower:
            col_map["ftti"] = col
        elif "注释" in header or "备注" in header:
            col_map["note"] = col

    print(f"\n列映射结果:")
    for key, col in sorted(col_map.items()):
        print(f"  {key}: 第{col}列 ({headers.get(col, 'N/A')})")

    # 数据起始行
    data_start_row = header_row + 1

    # 提取数据
    records = []
    current_function = None  # 用于跟踪合并单元格的功能名
    current_failure_id = None  # 用于跟踪合并单元格的failure_id
    current_hazard_id = None  # 用于跟踪合并单元格的hazard_id

    for row in range(data_start_row, ws.max_row + 1):
        # 检查是否是有效数据行（至少有异常表现或危害描述）
        anomaly = get_cell_value(ws, row, col_map.get("anomaly", 4))
        vehicle_hazard = get_cell_value(ws, row, col_map.get("vehicle_hazard", 5))

        if not anomaly and not vehicle_hazard:
            continue

        # 提取所有字段
        hazard_id = get_cell_value(ws, row, col_map.get("hazard_id", 1))
        function = get_cell_value(ws, row, col_map.get("function", 2))
        failure_id = get_cell_value(ws, row, col_map.get("failure_id", 3))

        # 处理合并单元格：如果当前单元格为空，使用上一个非空值
        if hazard_id:
            current_hazard_id = str(hazard_id).strip()
        else:
            hazard_id = current_hazard_id

        if function:
            current_function = normalize_text(function)
        else:
            function = current_function

        if failure_id:
            current_failure_id = str(failure_id).strip()
        else:
            failure_id = current_failure_id

        record = {
            "hazard_id": hazard_id if hazard_id else "",
            "function": function if function else "",
            "failure_id": failure_id if failure_id else "",
            "anomaly": normalize_text(anomaly) if anomaly else "",
            "vehicle_hazard": normalize_text(vehicle_hazard) if vehicle_hazard else "",
            "scenario": normalize_text(get_cell_value(ws, row, col_map.get("scenario", 6))),
            "description": normalize_text(get_cell_value(ws, row, col_map.get("description", 7))),
        }

        # 提取SEC值
        def parse_sec(val):
            if val is None:
                return None
            val_str = str(val).strip()
            if val_str in ("", "-", "N/A", "n/a", "NA"):
                return None
            try:
                return int(float(val_str))
            except:
                return None

        severity_val = get_cell_value(ws, row, col_map.get("severity", 8))
        exposure_val = get_cell_value(ws, row, col_map.get("exposure", 10))
        controllability_val = get_cell_value(ws, row, col_map.get("controllability", 12))

        record["severity"] = parse_sec(severity_val)
        record["severity_reason"] = normalize_text(get_cell_value(ws, row, col_map.get("severity_reason", 9)))
        record["exposure"] = parse_sec(exposure_val)
        record["exposure_reason"] = normalize_text(get_cell_value(ws, row, col_map.get("exposure_reason", 11)))
        record["controllability"] = parse_sec(controllability_val)
        record["controllability_reason"] = normalize_text(get_cell_value(ws, row, col_map.get("controllability_reason", 13)))

        # ASIL
        asil_val = get_cell_value(ws, row, col_map.get("asil", 14))
        if asil_val:
            asil_str = str(asil_val).strip().upper()
            if asil_str and asil_str not in ("", "-", "N/A", "NA"):
                record["asil"] = asil_str
            else:
                record["asil"] = None
        else:
            record["asil"] = None

        # 安全目标相关
        sg_id = get_cell_value(ws, row, col_map.get("safety_goal_id", 15))
        record["safety_goal_id"] = str(sg_id).strip() if sg_id else ""
        record["safety_goal"] = normalize_text(get_cell_value(ws, row, col_map.get("safety_goal", 16)))
        record["safe_state"] = normalize_text(get_cell_value(ws, row, col_map.get("safe_state", 17)))
        record["ftti"] = normalize_text(get_cell_value(ws, row, col_map.get("ftti", 18)))
        record["note"] = normalize_text(get_cell_value(ws, row, col_map.get("note", 19)))

        records.append(record)

        # 打印前5条记录用于调试
        if len(records) <= 5:
            print(f"\n记录{len(records)} (第{row}行):")
            print(f"  hazard_id: {record['hazard_id']}")
            print(f"  function: {record['function']}")
            print(f"  failure_id: {record['failure_id']}")
            print(f"  anomaly: {record['anomaly'][:50]}...")
            print(f"  S={record['severity']}, E={record['exposure']}, C={record['controllability']}, ASIL={record['asil']}")

    return records

def build_event_matrix(hara_records):
    """
    构建类似CS域的event_matrix结构
    """
    # 提取唯一的场景
    scenarios = {}
    scenario_id_counter = 1
    for rec in hara_records:
        scenario_text = rec["scenario"]
        if scenario_text and scenario_text not in scenarios:
            scenarios[scenario_text] = f"PT_SC_{scenario_id_counter:02d}"
            scenario_id_counter += 1

    print(f"\n提取到 {len(scenarios)} 个唯一场景")

    # 提取唯一的功能
    functions = {}
    for rec in hara_records:
        func_text = rec["function"]
        if func_text and func_text not in functions:
            functions[func_text] = f"PT_FUNC_{len(functions)+1:02d}"

    print(f"提取到 {len(functions)} 个唯一功能")

    # 构建event_matrix
    event_matrix = []
    for rec in hara_records:
        scenario_text = rec["scenario"]
        scenario_id = scenarios.get(scenario_text, "PT_SC_99")

        event = {
            "analysis_unit_id": "PT_MAIN",
            "failure_id": rec["failure_id"] if rec["failure_id"] else "PT_MF_UNKNOWN",
            "scenario_id": scenario_id,
            "description": rec["description"],
            "severity": rec["severity"],
            "severity_reason": rec["severity_reason"],
            "exposure": rec["exposure"],
            "exposure_reason": rec["exposure_reason"],
            "controllability": rec["controllability"],
            "controllability_reason": rec["controllability_reason"],
            "expected_asil": rec["asil"],
            "safety_goal": rec["safety_goal"] if rec["safety_goal"] else None,
            "safe_state": rec["safe_state"] if rec["safe_state"] else None,
            "ftti": rec["ftti"] if rec["ftti"] else None,
            "vehicle_hazard_id": rec["hazard_id"] if rec["hazard_id"] else "PT_HZ_UNKNOWN",
        }
        event_matrix.append(event)

    return event_matrix, scenarios, functions

def generate_domain_pack(event_matrix, scenarios, functions):
    """
    生成Domain Pack JSON格式
    """
    domain_pack = {
        "schema_version": "1.0",
        "domain": "PT",
        "domain_name": "动力域",
        "status": "pt_risk_matrix_extracted",
        "provenance": {
            "migration_stage": "P0-1",
            "risk_matrix_source": "FUSA HARA PT DOMAIN_V1.1_山子高科.xlsx",
            "risk_matrix_source_policy": "从参考基线提取完整风险矩阵，正确处理合并单元格",
            "extraction_date": "2026-08-26"
        },
        "function_catalog": {
            "function_matchers": [],
            "semantic_roles": {},
            "scope_rules": []
        },
        "analysis_catalog": {
            "analysis_units": [],
            "failure_mode_rules": {},
            "hazop_patterns": []
        },
        "risk_catalog": {
            "scenario_sets": {
                "pt_main": {
                    "scenarios": [
                        {"scenario_id": sid, "text": stext}
                        for stext, sid in sorted(scenarios.items(), key=lambda x: x[1])
                    ]
                }
            },
            "scenario_lookup": {},
            "event_matrix": event_matrix
        },
        "safety_goal_catalog": {}
    }

    return domain_pack

def main():
    # 设置输出编码
    sys.stdout.reconfigure(encoding='utf-8')

    # 读取Excel文件
    excel_path = r"C:\Users\zhouguangjian\Downloads\FUSA HARA PT DOMAIN_V1.1_山子高科 - 副本.xlsx"

    print(f"读取Excel文件: {excel_path}")
    wb = openpyxl.load_workbook(excel_path, data_only=True)

    print(f"\n工作表列表: {wb.sheetnames}")

    # 查找HARA分析表
    hara_sheet = None
    for name in wb.sheetnames:
        if "HARA" in name or "hara" in name:
            hara_sheet = wb[name]
            print(f"\n使用HARA表: '{name}'")
            break

    if not hara_sheet:
        print("错误：未找到HARA分析表")
        return

    # 提取数据
    print("\n" + "="*60)
    print("开始提取HARA数据")
    print("="*60)
    hara_records = extract_hara_sheet(hara_sheet)
    print(f"\n总共提取到 {len(hara_records)} 条HARA记录")

    # 构建event_matrix
    print("\n" + "="*60)
    print("构建event_matrix")
    print("="*60)
    event_matrix, scenarios, functions = build_event_matrix(hara_records)
    print(f"event_matrix包含 {len(event_matrix)} 条记录")

    # 生成Domain Pack
    print("\n" + "="*60)
    print("生成Domain Pack")
    print("="*60)
    domain_pack = generate_domain_pack(event_matrix, scenarios, functions)

    # 保存
    output_path = Path("references/domain_packs/PT_from_baseline.json")
    output_path.parent.mkdir(parents=True, exist_ok=True)

    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(domain_pack, f, ensure_ascii=False, indent=2)

    print(f"\nDomain Pack已保存到: {output_path}")

    # 统计信息
    print("\n" + "="*60)
    print("统计信息")
    print("="*60)
    print(f"场景数量: {len(scenarios)}")
    print(f"功能数量: {len(functions)}")
    print(f"事件矩阵记录数: {len(event_matrix)}")

    # ASIL分布
    asil_dist = defaultdict(int)
    for event in event_matrix:
        asil = event.get("expected_asil")
        asil_dist[asil if asil else "null"] += 1

    print("\nASIL分布:")
    for asil, count in sorted(asil_dist.items(), key=lambda x: str(x[0])):
        print(f"  {asil}: {count}")

    # S值分布
    s_dist = defaultdict(int)
    for event in event_matrix:
        s = event.get("severity")
        s_dist[str(s)] += 1

    print("\nS值分布:")
    for s, count in sorted(s_dist.items()):
        print(f"  S={s}: {count}")

    # 输出前3条完整样例
    print("\n" + "="*60)
    print("样例数据（前3条）")
    print("="*60)
    for i, event in enumerate(event_matrix[:3], 1):
        print(f"\n--- 记录 {i} ---")
        print(f"vehicle_hazard_id: {event['vehicle_hazard_id']}")
        print(f"failure_id: {event['failure_id']}")
        print(f"scenario_id: {event['scenario_id']}")
        print(f"description: {event['description'][:80]}...")
        print(f"S={event['severity']}, E={event['exposure']}, C={event['controllability']}")
        print(f"ASIL: {event['expected_asil']}")
        print(f"safety_goal: {event['safety_goal'][:60] if event['safety_goal'] else 'None'}...")

if __name__ == "__main__":
    main()
