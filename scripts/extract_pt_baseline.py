#!/usr/bin/env python3
"""
从PT域参考Excel提取数据并生成Domain Pack格式的预填表
"""

import json
import openpyxl
from pathlib import Path
from collections import defaultdict
import re

def normalize_text(text):
    """标准化文本：去除多余空格和换行"""
    if not text or not isinstance(text, str):
        return ""
    return " ".join(text.strip().split())

def extract_hara_sheet(ws):
    """
    提取HARA分析表数据
    表头在第2行
    """
    # 读取表头
    headers = {}
    for col in range(1, ws.max_column + 1):
        cell_value = ws.cell(2, col).value
        if cell_value:
            headers[col] = str(cell_value).strip()

    print(f"HARA表头: {headers}")

    # 定位关键列
    col_map = {}
    for col, header in headers.items():
        if "危害事件" in header and "ID" in header:
            col_map["hazard_id"] = col
        elif "整车功能" in header and "ID" not in header:
            col_map["function"] = col
        elif "功能失效" in header and "ID" in header:
            col_map["failure_id"] = col
        elif "功能异常表现" in header:
            col_map["anomaly"] = col
        elif "整车危害" in header:
            col_map["vehicle_hazard"] = col
        elif "运行场景" in header:
            col_map["scenario"] = col
        elif "危害事件描述" in header:
            col_map["description"] = col
        elif header.startswith("严重度") and "S" in header:
            col_map["severity"] = col
        elif "理由" in header and "S" in header:
            col_map["severity_reason"] = col
        elif "暴露概率" in header and "E" in header:
            col_map["exposure"] = col
        elif "理由" in header and "E" in header:
            col_map["exposure_reason"] = col
        elif "可控性" in header and "C" in header:
            col_map["controllability"] = col
        elif "理由" in header and "C" in header:
            col_map["controllability_reason"] = col
        elif header == "ASIL":
            col_map["asil"] = col
        elif "安全目标ID" in header:
            col_map["safety_goal_id"] = col
        elif "安全目标" in header and "ID" not in header:
            col_map["safety_goal"] = col
        elif "安全状态" in header:
            col_map["safe_state"] = col
        elif "FTTI" in header:
            col_map["ftti"] = col
        elif "注释" in header:
            col_map["note"] = col

    print(f"列映射: {col_map}")

    # 提取数据（从第3行开始）
    records = []
    for row in range(3, ws.max_row + 1):
        # 检查是否是空行
        hazard_id = ws.cell(row, col_map.get("hazard_id", 1)).value
        if not hazard_id or str(hazard_id).strip() == "":
            continue

        # 提取整行数据
        record = {
            "hazard_id": str(hazard_id).strip() if hazard_id else "",
            "function": normalize_text(ws.cell(row, col_map.get("function", 2)).value),
            "failure_id": str(ws.cell(row, col_map.get("failure_id", 3)).value or "").strip(),
            "anomaly": normalize_text(ws.cell(row, col_map.get("anomaly", 4)).value),
            "vehicle_hazard": normalize_text(ws.cell(row, col_map.get("vehicle_hazard", 5)).value),
            "scenario": normalize_text(ws.cell(row, col_map.get("scenario", 6)).value),
            "description": normalize_text(ws.cell(row, col_map.get("description", 7)).value),
        }

        # 提取SEC值（可能是数字或字符串）
        severity_val = ws.cell(row, col_map.get("severity", 8)).value
        exposure_val = ws.cell(row, col_map.get("exposure", 10)).value
        controllability_val = ws.cell(row, col_map.get("controllability", 12)).value

        # 转换为整数或null
        def parse_sec(val):
            if val is None or str(val).strip() in ("", "-", "N/A", "n/a"):
                return None
            try:
                return int(float(val))
            except:
                return None

        record["severity"] = parse_sec(severity_val)
        record["severity_reason"] = normalize_text(ws.cell(row, col_map.get("severity_reason", 9)).value)
        record["exposure"] = parse_sec(exposure_val)
        record["exposure_reason"] = normalize_text(ws.cell(row, col_map.get("exposure_reason", 11)).value)
        record["controllability"] = parse_sec(controllability_val)
        record["controllability_reason"] = normalize_text(ws.cell(row, col_map.get("controllability_reason", 13)).value)

        # ASIL
        asil_val = ws.cell(row, col_map.get("asil", 14)).value
        if asil_val and str(asil_val).strip() not in ("", "-", "N/A"):
            record["asil"] = str(asil_val).strip()
        else:
            record["asil"] = None

        # 安全目标相关
        record["safety_goal_id"] = str(ws.cell(row, col_map.get("safety_goal_id", 15)).value or "").strip()
        record["safety_goal"] = normalize_text(ws.cell(row, col_map.get("safety_goal", 16)).value)
        record["safe_state"] = normalize_text(ws.cell(row, col_map.get("safe_state", 17)).value)
        record["ftti"] = normalize_text(ws.cell(row, col_map.get("ftti", 18)).value)
        record["note"] = normalize_text(ws.cell(row, col_map.get("note", 19)).value)

        records.append(record)

    return records

def extract_hazop_sheet(ws):
    """
    提取HAZOP分析表数据
    表头在第2行
    """
    headers = {}
    for col in range(1, ws.max_column + 1):
        cell_value = ws.cell(2, col).value
        if cell_value:
            headers[col] = str(cell_value).strip()

    print(f"HAZOP表头: {headers}")

    # 定位关键列
    col_map = {}
    for col, header in headers.items():
        if "整车功能" in header and "ID" in header:
            col_map["function_id"] = col
        elif "整车功能" in header and "ID" not in header:
            col_map["function"] = col
        elif "功能失效模式" in header:
            col_map["failure_mode"] = col
        elif "功能异常表现" in header:
            col_map["anomaly"] = col
        elif "功能失效" in header and "ID" in header:
            col_map["failure_id"] = col
        elif "整车危害" in header:
            col_map["vehicle_hazard"] = col
        elif "关联" in header and "HARA" in header:
            col_map["associated_hara"] = col

    print(f"HAZOP列映射: {col_map}")

    records = []
    for row in range(3, ws.max_row + 1):
        function_id = ws.cell(row, col_map.get("function_id", 1)).value
        if not function_id or str(function_id).strip() == "":
            continue

        record = {
            "function_id": str(function_id).strip(),
            "function": normalize_text(ws.cell(row, col_map.get("function", 2)).value),
            "failure_mode": normalize_text(ws.cell(row, col_map.get("failure_mode", 3)).value),
            "anomaly": normalize_text(ws.cell(row, col_map.get("anomaly", 4)).value),
            "failure_id": str(ws.cell(row, col_map.get("failure_id", 5)).value or "").strip(),
            "vehicle_hazard": normalize_text(ws.cell(row, col_map.get("vehicle_hazard", 6)).value),
            "associated_hara": str(ws.cell(row, col_map.get("associated_hara", 7)).value or "").strip(),
        }
        records.append(record)

    return records

def build_event_matrix(hara_records, hazop_records):
    """
    构建类似CS域的event_matrix结构
    """
    # 先从HAZOP建立功能-失效模式-异常的映射
    hazop_map = {}
    for h in hazop_records:
        key = (h["function"], h["failure_mode"])
        if key not in hazop_map:
            hazop_map[key] = []
        hazop_map[key].append(h)

    # 提取唯一的场景
    scenarios = {}
    scenario_id_counter = 1
    for rec in hara_records:
        scenario_text = rec["scenario"]
        if scenario_text and scenario_text not in scenarios:
            scenarios[scenario_text] = f"PT_SC_{scenario_id_counter:02d}"
            scenario_id_counter += 1

    print(f"提取到 {len(scenarios)} 个唯一场景")

    # 构建event_matrix
    event_matrix = []
    for rec in hara_records:
        scenario_text = rec["scenario"]
        scenario_id = scenarios.get(scenario_text, "PT_SC_99")

        event = {
            "analysis_unit_id": "PT_MAIN",  # 暂用通用ID
            "failure_id": rec["failure_id"],
            "scenario_id": scenario_id,
            "description": rec["description"],
            "severity": rec["severity"],
            "severity_reason": rec["severity_reason"],
            "exposure": rec["exposure"],
            "exposure_reason": rec["exposure_reason"],
            "controllability": rec["controllability"],
            "controllability_reason": rec["controllability_reason"],
            "expected_asil": rec["asil"],
            "safety_goal": rec["safety_goal"] if rec["safety_goal"] else None,
            "safe_state": rec["safe_state"] if rec["safe_state"] else None,
            "ftti": rec["ftti"] if rec["ftti"] else None,
            "vehicle_hazard_id": rec["hazard_id"],
        }
        event_matrix.append(event)

    return event_matrix, scenarios

def generate_domain_pack(event_matrix, scenarios):
    """
    生成Domain Pack JSON格式
    """
    domain_pack = {
        "schema_version": "1.0",
        "domain": "PT",
        "domain_name": "动力域",
        "status": "pt_risk_matrix_extracted",
        "provenance": {
            "migration_stage": "P0-1",
            "risk_matrix_source": "FUSA HARA PT DOMAIN_V1.1_山子高科.xlsx",
            "risk_matrix_source_policy": "从参考基线提取完整风险矩阵",
            "extraction_date": "2026-08-26"
        },
        "function_catalog": {
            "function_matchers": [],
            "semantic_roles": {},
            "scope_rules": []
        },
        "analysis_catalog": {
            "analysis_units": [],
            "failure_mode_rules": {},
            "hazop_patterns": []
        },
        "risk_catalog": {
            "scenario_sets": {
                "pt_main": {
                    "scenarios": [
                        {"scenario_id": sid, "text": stext}
                        for stext, sid in sorted(scenarios.items(), key=lambda x: x[1])
                    ]
                }
            },
            "scenario_lookup": {},
            "event_matrix": event_matrix
        },
        "safety_goal_catalog": {}
    }

    return domain_pack

def main():
    # 读取Excel文件
    excel_path = r"C:\Users\zhouguangjian\Downloads\FUSA HARA PT DOMAIN_V1.1_山子高科 - 副本.xlsx"

    print(f"读取Excel文件: {excel_path}")
    wb = openpyxl.load_workbook(excel_path, data_only=True)

    print(f"工作表列表: {wb.sheetnames}")

    # 查找HARA分析表
    hara_sheet = None
    hazop_sheet = None

    for name in wb.sheetnames:
        if "HARA" in name or "hara" in name:
            hara_sheet = wb[name]
            print(f"找到HARA表: {name}")
        elif "HAZOP" in name or "hazop" in name:
            hazop_sheet = wb[name]
            print(f"找到HAZOP表: {name}")

    if not hara_sheet:
        print("错误：未找到HARA分析表")
        return

    # 提取数据
    print("\n=== 提取HARA数据 ===")
    hara_records = extract_hara_sheet(hara_sheet)
    print(f"提取到 {len(hara_records)} 条HARA记录")

    hazop_records = []
    if hazop_sheet:
        print("\n=== 提取HAZOP数据 ===")
        hazop_records = extract_hazop_sheet(hazop_sheet)
        print(f"提取到 {len(hazop_records)} 条HAZOP记录")

    # 构建event_matrix
    print("\n=== 构建event_matrix ===")
    event_matrix, scenarios = build_event_matrix(hara_records, hazop_records)
    print(f"event_matrix包含 {len(event_matrix)} 条记录")

    # 生成Domain Pack
    print("\n=== 生成Domain Pack ===")
    domain_pack = generate_domain_pack(event_matrix, scenarios)

    # 保存
    output_path = Path("references/domain_packs/PT_from_baseline.json")
    output_path.parent.mkdir(parents=True, exist_ok=True)

    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(domain_pack, f, ensure_ascii=False, indent=2)

    print(f"\nDomain Pack已保存到: {output_path}")

    # 统计信息
    print("\n=== 统计信息 ===")
    print(f"场景数量: {len(scenarios)}")
    print(f"事件矩阵记录数: {len(event_matrix)}")

    # ASIL分布
    asil_dist = defaultdict(int)
    for event in event_matrix:
        asil = event.get("expected_asil")
        asil_dist[asil if asil else "null"] += 1

    print("\nASIL分布:")
    for asil, count in sorted(asil_dist.items()):
        print(f"  {asil}: {count}")

    # S值分布
    s_dist = defaultdict(int)
    for event in event_matrix:
        s = event.get("severity")
        s_dist[s if s is not None else "null"] += 1

    print("\nS值分布:")
    for s, count in sorted(s_dist.items(), key=lambda x: (x[0] is None, x[0])):
        print(f"  S={s}: {count}")

    # 输出前3条样例
    print("\n=== 样例数据（前3条） ===")
    for i, event in enumerate(event_matrix[:3], 1):
        print(f"\n记录 {i}:")
        print(f"  vehicle_hazard_id: {event['vehicle_hazard_id']}")
        print(f"  failure_id: {event['failure_id']}")
        print(f"  scenario_id: {event['scenario_id']}")
        print(f"  S={event['severity']}, E={event['exposure']}, C={event['controllability']}")
        print(f"  ASIL: {event['expected_asil']}")
        print(f"  description: {event['description'][:60]}...")

if __name__ == "__main__":
    main()
