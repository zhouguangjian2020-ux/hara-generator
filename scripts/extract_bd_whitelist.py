"""从数据组11参考Excel提取BD域标准anomaly命名，生成bd_anomaly_whitelist.json"""
import json
from pathlib import Path
from openpyxl import load_workbook

REF_EXCEL = r"r:\Users\zhouguangjian\Developer\projects\数据组\数据组11\ALL-HARA_FUSA-500-GPFA-H002______FUSA HARA BD DOMAIN.xlsx"
OUTPUT = Path(__file__).parent.parent / "references" / "bd_anomaly_whitelist.json"


def main():
    wb = load_workbook(REF_EXCEL, read_only=True, data_only=True)
    ws = wb["HAZOP 分析"]

    rows = list(ws.iter_rows(values_only=True))

    # 表头在第2行（index=1）
    # col[0]: 整车功能 ID
    # col[1]: 整车功能
    # col[2]: 功能失效模式
    # col[3]: 功能异常表现
    # col[4]: 功能失效 ID
    # col[5]: 整车危害
    # col[6]: 关联 HARA

    # 分类映射（根据功能名归类）
    def categorize(func_name: str) -> str:
        f = func_name or ""
        if any(k in f for k in ["位置灯", "前照灯", "远光灯", "近光灯", "前雾灯", "后雾灯",
                                  "转向灯", "日行灯", "倒车灯", "制动灯", "牌照灯", "尾灯"]):
            return "外部灯光"
        if "雨刮" in f:
            return "雨刮"
        if "后视镜" in f:
            return "后视镜"
        if any(k in f for k in ["车窗", "天窗", "门窗"]):
            return "门窗天窗"
        if "门锁" in f or "儿童锁" in f:
            return "门锁"
        if "喇叭" in f or " horn" in f.lower():
            return "喇叭"
        if any(k in f for k in ["仪表", "报警", "指示"]):
            return "仪表报警"
        return "其他"

    anomaly_list = []
    seen = set()  # 去重 (功能异常表现 + failure_id)
    functions_set = set()  # 收集所有功能名
    modes_set = set()      # 收集所有失效模式

    # 数据从第3行开始（index=2）
    for row in rows[2:]:
        if not row:
            continue
        func_id = (row[0] or "").strip() if isinstance(row[0], str) else row[0]
        func_name = (row[1] or "").strip() if isinstance(row[1], str) else row[1]
        mode = (row[2] or "").strip() if isinstance(row[2], str) else row[2]
        anomaly = (row[3] or "").strip() if isinstance(row[3], str) else row[3]
        failure_id = (row[4] or "").strip() if isinstance(row[4], str) else row[4]
        hazard = (row[5] or "").strip() if isinstance(row[5], str) else row[5]

        # 跳过空行 / 非数据行
        if not func_id or not isinstance(func_id, str):
            continue
        if not func_id.startswith("B_"):
            continue
        if not anomaly:
            continue

        key = (anomaly, failure_id)
        if key in seen:
            continue
        seen.add(key)

        anomaly_list.append({
            "failure_id": failure_id,
            "function": func_name,
            "mode": mode,
            "anomaly": anomaly,
            "hazard": hazard,
            "category": categorize(func_name),
        })

        if func_name:
            functions_set.add(func_name)
        if mode:
            modes_set.add(mode)

    # BD 域标准失效模式（从"失效模式"sheet表头提取，丢失/非预期/间歇/过多/过少/过早）
    # 这些是车身功能的通用失效模式维度
    standard_modes = ["丢失", "非预期", "间歇", "过多", "过少", "过早"]
    wb.close()

    # 生成粗分类关键词：所有 {功能名}{标准失效模式} 的组合
    # 这些是 Agent 容易自创的粗分类命名（如"位置灯丢失"），应被识别并拒绝
    fine_anomalies = {a["anomaly"] for a in anomaly_list}
    coarse_keywords_set = set()
    for f in functions_set:
        for m in standard_modes:
            coarse = f + m
            # 仅当粗分类命名不与细分类命名重合时才加入
            if coarse not in fine_anomalies:
                coarse_keywords_set.add(coarse)
    coarse_keywords = sorted(coarse_keywords_set)

    wl = {
        "version": "1.0",
        "domain": "BD",
        "domain_name": "车身域",
        "description": "BD域车身功能的标准'功能异常表现'命名白名单，从数据组11参考Excel提取。"
                       "Agent生成S3 HAZOP时必须使用这些标准命名，不得自创粗分类（如'位置灯丢失'、'前雾灯非预期'等）。",
        "total_anomalies": len(anomaly_list),
        "coarse_keywords": coarse_keywords,
        "anomaly_list": anomaly_list,
    }

    # 按category统计
    cat_count = {}
    for a in anomaly_list:
        cat = a["category"]
        cat_count[cat] = cat_count.get(cat, 0) + 1
    print(f"Total anomalies: {len(anomaly_list)}")
    print(f"Categories: {json.dumps(cat_count, ensure_ascii=False, indent=2)}")
    print(f"Coarse keywords ({len(coarse_keywords)}): {coarse_keywords}")

    # 输出每个分类前几条
    by_cat = {}
    for a in anomaly_list:
        by_cat.setdefault(a["category"], []).append(a)
    for cat, items in by_cat.items():
        print(f"\n=== {cat} ({len(items)} 条) ===")
        for it in items[:3]:
            print(f"  [{it['failure_id']}] {it['function']} | {it['mode']} | {it['anomaly']}")

    OUTPUT.write_text(json.dumps(wl, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"\nWritten to: {OUTPUT}")


if __name__ == "__main__":
    main()
